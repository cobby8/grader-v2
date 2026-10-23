# -*- coding: utf-8 -*-
"""재단 패턴(F10 1단계 «보기 전용») — 스타일·버전 창고 + /api/cad/* 라우터.

왜 이 파일이 있나(비유):
  stizshop 의 패턴 엔진이 미리 만든 «화면용 그림 자료(view.json)» 와 «공장용 산출물 묶음(exports.zip)» 을
  관리자가 올리면, 이 서버는 그것을 «사진첩» 처럼 보관하고 직원에게 보여 주기만 한다.
  서버에는 패턴 엔진이 없다(1단계 — 메모리·빌드 위험 0). 계산은 전부 stizshop 에서 끝나 있다.

저장 구조:
  디스크  data/cad/{style_id}/v{n}/view.json · exports.zip · meta.json   (+ backup.json = 이 서버의 백업 기록, 창고엔 안 올림)
  창고    Supabase Storage 비공개 창고 «pattern-cad»
            {style_id}/v{n}/view · exports · meta        ← 열쇠는 영문 소문자·숫자·-·/ 만 (07-09 한글 열쇠 무음 실패 사고)
  버전은 «덮어쓰지 않고 늘린다» — 같은 스타일을 다시 올리면 v2, v3 … 옛 버전은 그대로 남는다.

권한:
  🔒 보기·내려받기 = 로그인 직원 (main.py 가 /api 전역에 require_auth 를 건다)
  🛡 올리기 = 관리자 (admin_required)
  창고 쓰기·읽기는 «요청자 로그인 토큰을 그대로 넘긴다» — 비밀키를 쓰지 않는다(storage_backup.py 와 같은 원칙).
  창고 규칙: 읽기 = 로그인(authenticated) · 쓰기 = 관리자 · 덮어쓰기(UPDATE)·삭제 규칙 없음.

재배포 복원:
  창고 읽기가 «로그인 사용자만» 이라 부팅 때(토큰 없음)는 읽을 수 없다(비밀키 금지).
  그래서 «로그인한 사람이 처음 목록을 열 때» 그 사람 토큰으로 디스크에 없는 버전만 내려받아 복원한다(10분에 한 번).
"""
from __future__ import annotations

import hashlib
import importlib.util
import io
import json
import os
import re
import shutil
import sys
import tempfile
import threading
import time
import traceback
import zipfile
from urllib.parse import quote
from typing import Any, Dict, List, Optional, Tuple

import httpx
from fastapi import APIRouter, Body, Depends, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, JSONResponse, Response
from starlette.background import BackgroundTask
from starlette.concurrency import run_in_threadpool

from . import state
from .auth import admin_required, auth_required

router = APIRouter(prefix="/api/cad", tags=["cad"])

# ── 상수 ────────────────────────────────────────────────────────────────
BUCKET = "pattern-cad"
VIEW_FORMAT = "stiz-cad-view"
# ⚠️ 세 정규식은 반드시 fullmatch 로 쓴다 — `$` 는 끝의 줄바꿈 하나를 봐줘서 "abc\n" 이 통과한다(검증 리뷰 지적)
STYLE_ID_RE = re.compile(r"[a-z0-9][a-z0-9-]{1,62}[a-z0-9]")         # 영문 소문자·숫자·- 만(디스크 폴더·창고 열쇠에 그대로 쓴다)
KEY_RE = re.compile(r"[a-z0-9/-]+")                                  # 창고 열쇠 검사 — 한글·점·공백이 섞이면 올리지 않는다
ZIP_MEMBER_RE = re.compile(r"[A-Za-z0-9][A-Za-z0-9_./-]*")           # ZIP 안 이름 — 영문·숫자만(stizshop cad_view.py 가 그렇게 만든다)
MAX_VIEW_BYTES = 20 * 1024 * 1024
MAX_ZIP_BYTES = 50 * 1024 * 1024                                     # 창고 file_size_limit 와 같다
MAX_UNZIPPED_TOTAL = 200 * 1024 * 1024                               # ZIP 을 풀었을 때 합계 상한(압축 폭탄 방어)
MAX_UNZIPPED_FILE = 60 * 1024 * 1024                                 # ZIP 안 파일 하나 상한
MAX_ZIP_MEMBERS = 2000
OBJECTS = (("view.json", "view", "application/json"), ("exports.zip", "exports", "application/zip"),
           ("meta.json", "meta", "application/json"))               # (디스크 이름, 창고 이름, 종류)
# 있을 때만 올리고·복원하는 것(F10 3·5단계) — 규칙표(다시 그레이딩한 버전) · 잠금 기록(공장 보낸 판)
#   창고는 «덮어쓰기 금지» 라 meta 를 고쳐 다시 올릴 수 없다 → 잠금은 따로 «lock» 물건으로 남긴다
OPTIONAL_OBJECTS = (("rules.json", "rules", "application/json"), ("lock.json", "lock", "application/json"))
FILE_KINDS = {                                                       # 내려받기 kind → (보통 종류, 설명)
    "zip": ("application/zip", "산출물 전체 ZIP"),
    "dxf": ("application/dxf", "사이즈별 DXF"),
    "pdf": ("application/pdf", "사이즈별 1:1 PDF"),
    "ai": ("application/postscript", "사이즈별 .ai"),
    "aama": ("application/dxf", "AAMA 기준 사이즈 DXF"),
    "rul": ("text/plain; charset=utf-8", "AAMA 규칙표 RUL"),
    "allsizes": ("application/dxf", "전 사이즈 겹친 DXF"),
    "a4": ("application/pdf", "A4 분할 PDF"),
    # 들여온 스타일(F10 6단계 · stizshop style_import 묶음) — 기준 명세·규칙표·검산·확인 목록·원본
    "spec_base": ("application/json", "기준 사이즈 명세"),
    "spec": ("application/json", "사이즈별 명세"),
    "rules": ("application/json", "그레이딩 규칙표(JSON)"),
    "check": ("application/json", "원본 대조 검산"),
    "confirm": ("text/markdown; charset=utf-8", "확인 필요 목록"),
    "names": ("application/json", "이름 확인 자료"),
    "source": ("application/postscript", "원본 OptiCad .ai"),
}
META_EXTRA_KEYS = ("rebuilt_from", "engine_run", "regraded_from", "rule_changes", "mesh",
                   "imported", "renamed_from", "name_changes",       # (F10 6단계 들여오기·이름 확인)
                   "intake")                                         # 패턴 의뢰에서 자동으로 만든 판(2026-10-22) — 의뢰 번호·검토 대기   # store_version(extra=…) 로 더할 수 있는 칸
_HTTP_TIMEOUT = 30.0
_RESTORE_EVERY = 600.0                                               # 복원 확인 간격(초)

_LOCK = threading.Lock()                                             # 올리기·복원이 같은 폴더를 동시에 만지지 않게
_last_restore = {"at": 0.0, "result": None, "busy": False}           # 복원 시각·결과 — 읽고 쓰기는 _LOCK 안에서만


def _log(msg: str) -> None:
    """서버 콘솔에 한 줄 — 콘솔이 cp949(윈도우)여도 «−»·«→» 같은 글자 때문에 작업 스레드가 죽지 않게(못 쓰는 글자는 ? 로)."""
    try:
        print(msg, flush=True)
    except UnicodeEncodeError:
        enc = getattr(sys.stdout, "encoding", None) or "ascii"
        try:
            print(msg.encode(enc, errors="replace").decode(enc, errors="replace"), flush=True)
        except Exception:
            pass
    except Exception:
        pass


# ══════════════════════════════════════════════════════════════════════
#  디스크
# ══════════════════════════════════════════════════════════════════════
def cad_dir() -> str:
    """data/cad 폴더(환경변수 GRADER_CAD_DIR 로 바꿀 수 있다 — 시험용)."""
    d = os.environ.get("GRADER_CAD_DIR") or os.path.join(state.PROJECT_ROOT, "data", "cad")
    os.makedirs(d, exist_ok=True)
    return d


def _check_style_id(style_id: str) -> str:
    if not isinstance(style_id, str) or not STYLE_ID_RE.fullmatch(style_id):
        raise HTTPException(status_code=400, detail="스타일 번호는 영문 소문자·숫자·- 만 쓸 수 있습니다(3~64자).")
    return style_id


def _versions_on_disk(style_id: str) -> List[int]:
    d = os.path.join(cad_dir(), style_id)
    if not os.path.isdir(d):
        return []
    out = []
    for nm in os.listdir(d):
        m = re.fullmatch(r"v([1-9][0-9]{0,5})", nm)
        if m and os.path.isfile(os.path.join(d, nm, "meta.json")):
            out.append(int(m.group(1)))
    return sorted(out)


def _vdir(style_id: str, ver: int) -> str:
    return os.path.join(cad_dir(), style_id, f"v{ver}")


def _read_meta(style_id: str, ver: int) -> Optional[Dict[str, Any]]:
    try:
        with open(os.path.join(_vdir(style_id, ver), "meta.json"), encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return None


def _read_backup(style_id: str, ver: int) -> Optional[Dict[str, Any]]:
    try:
        with open(os.path.join(_vdir(style_id, ver), "backup.json"), encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return None


def _resolve_ver(style_id: str, ver: str) -> int:
    vs = _versions_on_disk(style_id)
    if not vs:
        raise HTTPException(status_code=404, detail="그 스타일이 없습니다.")
    if ver == "latest":
        return vs[-1]
    if not re.fullmatch(r"[1-9][0-9]{0,5}", ver or "") or int(ver) not in vs:
        raise HTTPException(status_code=404, detail="그 버전이 없습니다.")
    return int(ver)


def _sha256(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()


def _write_version(style_id: str, ver: int, blobs: Dict[str, bytes]) -> None:
    """임시 폴더에 다 쓴 뒤 이름만 바꿔 넣는다 — 반쯤 쓴 버전이 목록에 보이지 않게."""
    final = _vdir(style_id, ver)
    if os.path.exists(final):
        raise HTTPException(status_code=409, detail=f"v{ver} 이 이미 있습니다(덮어쓰지 않습니다).")
    tmp = final + f".tmp-{os.getpid()}-{int(time.time() * 1000)}"
    os.makedirs(tmp)
    try:
        for name, data in blobs.items():
            with open(os.path.join(tmp, name), "wb") as f:
                f.write(data)
        os.replace(tmp, final)
    except Exception:
        shutil.rmtree(tmp, ignore_errors=True)
        raise


# ══════════════════════════════════════════════════════════════════════
#  창고(Supabase Storage) — 요청자 토큰을 그대로 넘긴다(비밀키 없음)
# ══════════════════════════════════════════════════════════════════════
def _sb_config() -> Tuple[Optional[str], Optional[str]]:
    base = os.environ.get("SUPABASE_URL"); pub = os.environ.get("SUPABASE_PUBLISHABLE_KEY")
    if not base or not pub:
        return None, None
    return base.rstrip("/"), pub


def _bearer(request: Request) -> Optional[str]:
    authz = request.headers.get("Authorization") or request.headers.get("authorization") or ""
    if authz.lower().startswith("bearer "):
        tok = authz.split(" ", 1)[1].strip()
        return tok or None
    return None


def object_key(style_id: str, ver: int, name: str) -> str:
    """창고 열쇠 — 영문 소문자·숫자·-·/ 만. 어기면 예외(조용히 실패하지 않게)."""
    key = f"{style_id}/v{ver}/{name}"
    if not KEY_RE.fullmatch(key):
        raise ValueError(f"창고 열쇠에 쓸 수 없는 글자: {key!r}")
    return key


def _sb_headers(pub: str, jwt: str, ctype: Optional[str] = None) -> Dict[str, str]:
    h = {"apikey": pub, "Authorization": f"Bearer {jwt}"}
    if ctype:
        h["Content-Type"] = ctype
    return h


def _sb_upload(key: str, data: bytes, ctype: str, jwt: str) -> Tuple[bool, str]:
    base, pub = _sb_config()
    try:
        r = httpx.post(f"{base}/storage/v1/object/{BUCKET}/{key}", content=data,
                       headers=_sb_headers(pub, jwt, ctype), timeout=_HTTP_TIMEOUT)   # x-upsert 없음 = 덮어쓰기 금지
    except Exception as e:
        return False, f"네트워크 {type(e).__name__}"
    if r.status_code in (200, 201):
        return True, "ok"
    return False, f"status {r.status_code}"


def _sb_download(key: str, jwt: str) -> Optional[bytes]:
    base, pub = _sb_config()
    try:
        r = httpx.get(f"{base}/storage/v1/object/{BUCKET}/{key}", headers=_sb_headers(pub, jwt), timeout=_HTTP_TIMEOUT)
    except Exception:
        return None
    return r.content if r.status_code == 200 else None


def _sb_list(prefix: str, jwt: str) -> Optional[List[Dict[str, Any]]]:
    base, pub = _sb_config()
    try:
        r = httpx.post(f"{base}/storage/v1/object/list/{BUCKET}", headers=_sb_headers(pub, jwt, "application/json"),
                       json={"prefix": prefix, "limit": 1000, "offset": 0}, timeout=_HTTP_TIMEOUT)
    except Exception:
        return None
    if r.status_code != 200:
        return None
    try:
        items = r.json()
    except Exception:
        return None
    return items if isinstance(items, list) else None


def backup_version(style_id: str, ver: int, jwt: Optional[str]) -> Dict[str, Any]:
    """디스크의 한 버전을 창고에 올리고 «되읽어» 지문을 맞춰 본다. 실패해도 예외를 던지지 않는다."""
    base, pub = _sb_config()
    if not base:
        res = {"ok": False, "skipped": True, "reason": "창고 미설정(로컬)"}
    elif not jwt:
        res = {"ok": False, "skipped": True, "reason": "로그인 토큰 없음(로컬 무인증)"}
    else:
        res = {"ok": True, "objects": {}}
        for disk_name, store_name, ctype in OBJECTS:
            with open(os.path.join(_vdir(style_id, ver), disk_name), "rb") as f:
                data = f.read()
            key = object_key(style_id, ver, store_name)
            ok, why = _sb_upload(key, data, ctype, jwt)
            back = _sb_download(key, jwt) if ok else None                   # 되읽기 — 올라간 것이 정말 같은가
            verified = back is not None and _sha256(back) == _sha256(data)
            res["objects"][store_name] = {"key": key, "uploaded": ok, "verified": verified, "reason": why, "bytes": len(data)}
            if not verified:
                res["ok"] = False
        for disk_name, store_name, ctype in OPTIONAL_OBJECTS:               # 규칙표(다시 그레이딩한 버전)처럼 «있을 때만»
            if os.path.isfile(os.path.join(_vdir(style_id, ver), disk_name)):
                o = _backup_object(style_id, ver, disk_name, store_name, ctype, jwt)
                o.pop("existing", None)
                res["objects"][store_name] = o
                if not o.get("ok"):
                    res["ok"] = False
        if not res["ok"]:
            res["reason"] = "일부 파일이 창고에 올라가지 않았거나 되읽은 지문이 다릅니다"
    res["at"] = time.strftime("%Y-%m-%dT%H:%M:%S")
    try:
        with open(os.path.join(_vdir(style_id, ver), "backup.json"), "w", encoding="utf-8") as f:
            json.dump(res, f, ensure_ascii=False, indent=1)
    except Exception:
        pass
    tag = "성공" if res.get("ok") else ("건너뜀" if res.get("skipped") else "실패")
    _log(f"[cad] 창고 백업 {tag}: {style_id} v{ver}")                       # 토큰·키는 남기지 않는다
    return res


def restore_missing_with_token(jwt: Optional[str], force: bool = False) -> Dict[str, Any]:
    """창고에는 있는데 디스크에 없는 버전을 내려받는다(로그인 사용자 토큰으로). 10분에 한 번만.

    기존 storage_backup.restore_missing 과 같은 원칙: 디스크 우선(있는 것은 안 건드림) · 이상한 이름은 건너뜀 ·
    meta 의 지문과 다르면 버린다 · 예외를 밖으로 던지지 않는다.
    """
    base, _ = _sb_config()
    if not base or not jwt:
        return {"restored": 0, "skipped": True}
    with _LOCK:                                                             # 시각 확인과 «복원 중» 표시를 한 번에(두 요청이 겹쳐 돌지 않게)
        if _last_restore["busy"]:
            return {"restored": 0, "busy": True}
        if not force and time.time() - _last_restore["at"] < _RESTORE_EVERY:
            return {"restored": 0, "cached": True, "last": _last_restore["result"]}
        _last_restore["busy"] = True
    restored, errors, listed = 0, 0, False
    try:
        top = _sb_list("", jwt)
        if top is None:
            return {"restored": 0, "error": "창고 목록을 읽지 못했습니다"}      # 실패 — 시각을 남기지 않아 다음 요청이 다시 시도한다
        listed = True
        for it in top:
            sid = it.get("name") if isinstance(it, dict) else None
            if not isinstance(sid, str) or not STYLE_ID_RE.fullmatch(sid) or it.get("id") is not None:
                continue                                                    # 폴더(id 없음)만 · 이름 검사
            vers = _sb_list(sid + "/", jwt) or []
            for v in vers:
                vn = v.get("name") if isinstance(v, dict) else None
                m = re.fullmatch(r"v([1-9][0-9]{0,5})", vn or "")
                if not m or v.get("id") is not None:
                    continue
                ver = int(m.group(1))
                if os.path.isdir(_vdir(sid, ver)):
                    continue                                                # 디스크 우선
                try:
                    blobs = {}
                    for disk_name, store_name, _c in OBJECTS:
                        b = _sb_download(object_key(sid, ver, store_name), jwt)
                        if b is None:
                            raise ValueError("내려받기 실패")
                        blobs[disk_name] = b
                    meta = json.loads(blobs["meta.json"].decode("utf-8"))
                    if (meta.get("view_sha256") != _sha256(blobs["view.json"]) or
                            meta.get("zip_sha256") != _sha256(blobs["exports.zip"]) or meta.get("style_id") != sid
                            or meta.get("version") != ver):
                        raise ValueError("지문 불일치")
                    # 있을 때만인 것(규칙표·잠금) — 창고 목록에 «있다» 고 나오는데 못 받으면 이 버전은 다음에 다시(잠금이 조용히 풀리지 않게)
                    #   «목록을 못 읽음» 을 «없음» 으로 보면 잠금이 조용히 풀린다(검증 지적) → 그 버전은 복원하지 않고 다음에 다시
                    listed = _sb_list(f"{sid}/v{ver}/", jwt)
                    if listed is None:
                        raise ValueError("버전 폴더 목록을 읽지 못함(다음에 다시)")
                    listed_names = {(x.get("name") if isinstance(x, dict) else None) for x in listed}
                    if meta.get("rules_sha256") and "rules" not in listed_names:
                        raise ValueError("meta 에 규칙표 지문이 있는데 창고에 규칙표가 없음")
                    if meta.get("locked") and "lock" not in listed_names:
                        raise ValueError("meta 에 잠금 표시가 있는데 창고에 잠금 기록이 없음")
                    for disk_name, store_name, _c in OPTIONAL_OBJECTS:
                        if store_name not in listed_names:
                            continue
                        if store_name == "rules" and not meta.get("rules_sha256"):
                            continue                                        # meta 가 보증하지 않는 규칙표는 받지 않는다
                        b = _sb_download(object_key(sid, ver, store_name), jwt)
                        if b is None:
                            raise ValueError(f"{store_name} 내려받기 실패")
                        if store_name == "rules" and meta["rules_sha256"] != _sha256(b):
                            raise ValueError("규칙표 지문 불일치")
                        if store_name == "lock":
                            rec = json.loads(b.decode("utf-8"))
                            if not isinstance(rec, dict) or rec.get("style_id") != sid or rec.get("version") != ver:
                                raise ValueError("잠금 기록 불일치")
                        blobs[disk_name] = b
                    with _LOCK:
                        if not os.path.isdir(_vdir(sid, ver)):
                            _write_version(sid, ver, blobs)
                            with open(os.path.join(_vdir(sid, ver), "backup.json"), "w", encoding="utf-8") as f:
                                json.dump({"ok": True, "restored": True, "at": time.strftime("%Y-%m-%dT%H:%M:%S")}, f)
                            restored += 1
                            _log(f"[cad] 창고에서 복원: {sid} v{ver}")
                except Exception as e:
                    errors += 1
                    _log(f"[cad] 복원 건너뜀: {sid} v{ver} - {type(e).__name__}")
    except Exception as e:
        listed = False
        _log(f"[cad] 복원 중 오류(무시): {type(e).__name__}")
    finally:
        with _LOCK:
            _last_restore["busy"] = False
            if listed:                                                      # 목록을 끝까지 읽었을 때만 «확인함» 으로 기록
                _last_restore["at"] = time.time()
                _last_restore["result"] = {"restored": restored, "errors": errors}
    return {"restored": restored, "errors": errors}


# ══════════════════════════════════════════════════════════════════════
#  올린 파일 검사
# ══════════════════════════════════════════════════════════════════════
def _is_num(v: Any) -> bool:
    """진짜 숫자인가(참·거짓·NaN·무한대 제외)"""
    return isinstance(v, (int, float)) and not isinstance(v, bool) and v == v and abs(v) != float("inf")


def validate_upload(view_bytes: bytes, zip_bytes: bytes) -> Dict[str, Any]:
    """view.json + exports.zip 가 서로 맞는지 본다. 틀리면 400(한국어 사유).

    모양이 이상한 자료(글자 자리에 목록 등)도 500 이 아니라 400 으로 — 마지막에 형식 오류를 통째로 잡는다.
    """
    try:
        return _validate_upload(view_bytes, zip_bytes)
    except HTTPException:
        raise
    except (TypeError, AttributeError, KeyError, ValueError, IndexError) as e:
        raise HTTPException(status_code=400, detail=f"view.json 모양이 예상과 다릅니다({type(e).__name__}).")


def _validate_upload(view_bytes: bytes, zip_bytes: bytes) -> Dict[str, Any]:
    def bad(msg: str):
        raise HTTPException(status_code=400, detail=msg)
    if not view_bytes or len(view_bytes) > MAX_VIEW_BYTES:
        bad("view.json 이 비었거나 너무 큽니다(20MB 이하).")
    if not zip_bytes or len(zip_bytes) > MAX_ZIP_BYTES:
        bad("exports.zip 이 비었거나 너무 큽니다(50MB 이하).")
    try:
        view = json.loads(view_bytes.decode("utf-8"))
    except Exception:
        bad("view.json 을 읽지 못했습니다(JSON 형식이 아님).")
    if not isinstance(view, dict) or view.get("format") != VIEW_FORMAT:
        bad("view.json 형식이 다릅니다(stizshop tools/pattern-engine/cad_view.py 결과만 받습니다).")
    style = view.get("style")
    if not isinstance(style, dict):
        bad("view.json 의 style 은 객체여야 합니다.")
    sid = style.get("id")
    if not isinstance(sid, str) or not STYLE_ID_RE.fullmatch(sid):
        bad("view.json 의 스타일 번호(style.id)는 영문 소문자·숫자·- 만 쓸 수 있습니다.")
    for k in ("name", "garment", "base_size", "graded_from"):
        if style.get(k) is not None and not isinstance(style.get(k), str):
            bad(f"view.json 의 style.{k} 는 글자여야 합니다.")
    sizes = style.get("sizes")
    if not isinstance(sizes, list) or not sizes or not all(isinstance(x, str) and x for x in sizes):
        bad("view.json 의 style.sizes 는 사이즈 이름 목록이어야 합니다.")
    if not isinstance(view.get("sizes"), dict) or any(not isinstance(view["sizes"].get(x), dict) for x in sizes):
        bad("view.json 에 사이즈 자료가 빠져 있습니다.")
    for x in sizes:
        if not isinstance(view["sizes"][x].get("pieces"), list) or not isinstance(view["sizes"][x].get("checks", []), list):
            bad(f"view.json 사이즈 {x[:20]} 의 pieces·checks 는 목록이어야 합니다.")
    if not isinstance(style.get("pieces", []), list):
        bad("view.json 의 style.pieces 는 목록이어야 합니다.")
    # 화면이 숫자로 쓰는 값 — 숫자가 아니면 받지 않는다(글자가 화면 HTML 에 섞이지 않게)
    if not _is_num(view.get("q")) or view["q"] <= 0:
        bad("view.json 의 q(좌표 단위)는 양수여야 합니다.")
    for k in ("tol_check", "simplify"):
        if k in view and not (_is_num(view[k]) and view[k] >= 0):
            bad(f"view.json 의 {k} 는 숫자여야 합니다.")
    ez = view.get("exports_zip", {})
    if not isinstance(ez, dict):
        bad("view.json 의 exports_zip 은 객체여야 합니다.")
    if ez.get("sha256") is not None and (not isinstance(ez["sha256"], str) or ez["sha256"] != _sha256(zip_bytes)):
        bad("exports.zip 이 view.json 과 짝이 아닙니다(지문이 다름). 같은 때 만든 두 파일을 올려 주세요.")
    files = view.get("files", [])
    if not isinstance(files, list) or not all(isinstance(f, dict) for f in files):
        bad("view.json 의 files 는 객체 목록이어야 합니다.")
    for f in files:
        if not isinstance(f.get("path"), str) or not isinstance(f.get("sha256"), str) or not isinstance(f.get("kind"), str):
            bad("view.json 의 files 항목(path·sha256·kind)은 글자여야 합니다.")
        if f.get("size") is not None and not isinstance(f.get("size"), str):
            bad("view.json 의 files 항목 size 는 글자여야 합니다.")
        if f["kind"] not in FILE_KINDS:
            bad(f"모르는 파일 종류: {f['kind'][:20]}")
    try:
        zf = zipfile.ZipFile(io.BytesIO(zip_bytes))
    except zipfile.BadZipFile:
        bad("exports.zip 이 ZIP 이 아닙니다.")
    with zf:
        infos = zf.infolist()
        # 압축 폭탄 방어 — 풀기 전에 «풀면 몇 바이트인가» 를 먼저 본다(작은 ZIP 이 수 GB 로 부푸는 것 차단)
        if len(infos) > MAX_ZIP_MEMBERS:
            bad(f"ZIP 안 파일이 너무 많습니다({MAX_ZIP_MEMBERS}개 이하).")
        if any(i.file_size > MAX_UNZIPPED_FILE for i in infos):
            bad(f"ZIP 안에 너무 큰 파일이 있습니다(파일 하나 {MAX_UNZIPPED_FILE // 2**20}MB 이하).")
        if sum(i.file_size for i in infos) > MAX_UNZIPPED_TOTAL:
            bad(f"ZIP 을 풀면 너무 큽니다(합계 {MAX_UNZIPPED_TOTAL // 2**20}MB 이하).")
        names = {i.filename for i in infos}
        for n in names:
            if not ZIP_MEMBER_RE.fullmatch(n) or ".." in n.split("/"):
                bad(f"ZIP 안 파일 이름에 쓸 수 없는 글자가 있습니다: {n[:80]}")
        for f in files:
            p = f["path"]
            if p not in names:
                bad(f"ZIP 에 빠진 파일: {p[:80]}")
            if _sha256(zf.read(p)) != f["sha256"]:
                bad(f"ZIP 안 파일 지문이 view.json 과 다릅니다: {p[:80]}")
    return view


# ══════════════════════════════════════════════════════════════════════
#  API
# ══════════════════════════════════════════════════════════════════════
def _style_summary(style_id: str) -> Optional[Dict[str, Any]]:
    vs = _versions_on_disk(style_id)
    if not vs:
        return None
    metas = [m for m in (_read_meta(style_id, v) for v in vs) if m]
    if not metas:
        return None
    last = metas[-1]
    imp = last.get("imported") if isinstance(last.get("imported"), dict) else None
    return {"id": style_id, "name": last.get("name"), "garment": last.get("garment"), "sizes": last.get("sizes"),
            "imported": bool(imp), "confirm_open": (imp or {}).get("confirm_open"),       # 목록 카드의 «들여옴 · 확인 필요 N»
            "versions": vs, "latest": _public_meta(last, _read_backup(style_id, last.get("version"))),
            "locked_versions": [v for v in vs if _lock_info(style_id, v)]}          # 목록 카드의 «공장 보냄» 표시


def _public_meta(m: Dict[str, Any], backup: Optional[Dict[str, Any]]) -> Dict[str, Any]:
    out = {k: m.get(k) for k in ("version", "uploaded_at", "uploaded_by", "note", "sizes", "size_labels", "files",
                                 "view_sha256", "zip_sha256", "view_bytes", "zip_bytes", "source", "rebuilt_from", "engine_run",
                                 "regraded_from", "mesh", "imported", "renamed_from")}
    sid, ver = m.get("style_id"), m.get("version")
    out["locked"] = _lock_info(sid, ver) if (isinstance(sid, str) and isinstance(ver, int)) else None   # 공장 보낸 판(누가·언제)
    out["rule_changes"] = (m.get("rule_changes") or [])[:50]                 # 화면 표시용 앞부분만(전체는 rules.json)
    out["rule_changes_count"] = len(m.get("rule_changes") or [])
    out["name_changes"] = (m.get("name_changes") or [])[:50]                 # 이름 확인으로 바꾼 것(들여온 스타일)
    out["name_changes_count"] = len(m.get("name_changes") or [])
    out["has_rules"] = bool(m.get("rules_sha256"))
    out["backup"] = None if backup is None else {"ok": bool(backup.get("ok")), "skipped": bool(backup.get("skipped")),
                                                 "restored": bool(backup.get("restored")), "reason": backup.get("reason"),
                                                 "at": backup.get("at")}
    return out


def _can_upload(request: Request) -> bool:
    return (not auth_required()) or getattr(request.state, "role", None) == "admin"


@router.get("/me")
def cad_me(request: Request) -> Dict[str, Any]:
    """화면이 «올리기» 단추를 보여 줄지 묻는다(관리자이거나 로컬 무인증)."""
    base, _ = _sb_config()
    try:
        rebuildable = sorted(getattr(load_service(), "REBUILDABLE", {}))     # 서버에서 «다시 만들기» 가 되는 스타일
    except Exception:
        rebuildable = []
    try:
        markerable = rebuildable if hasattr(load_service(), "marker") else []   # 마커(F10 4단계)도 서버에 원본이 있는 스타일만
    except Exception:
        markerable = []
    try:
        importable = _can_upload(request) and hasattr(load_service(), "import_style")   # 새 스타일 들여오기(F10 6단계) — 관리자만
    except Exception:
        importable = False
    return {"can_upload": _can_upload(request), "auth_required": auth_required(), "storage": bool(base),
            "rebuildable": rebuildable, "markerable": markerable, "importable": importable}


@router.get("/styles")
def list_styles(request: Request) -> Dict[str, Any]:
    restore = restore_missing_with_token(_bearer(request))                 # 재배포 뒤 첫 로그인 사용자가 복원을 일으킨다
    out = []
    root = cad_dir()
    for sid in sorted(os.listdir(root)):
        if STYLE_ID_RE.fullmatch(sid) and os.path.isdir(os.path.join(root, sid)):
            s = _style_summary(sid)
            if s:
                out.append(s)
    return {"styles": out, "restore": restore}


@router.get("/styles/{style_id}")
def get_style(style_id: str, request: Request) -> Dict[str, Any]:
    _check_style_id(style_id)
    if not _versions_on_disk(style_id):
        restore_missing_with_token(_bearer(request))
    vs = _versions_on_disk(style_id)
    if not vs:
        raise HTTPException(status_code=404, detail="그 스타일이 없습니다.")
    s = _style_summary(style_id)
    if s is None:                                                           # meta.json 이 깨졌다 — 없는 것으로 본다
        raise HTTPException(status_code=404, detail="그 스타일의 정보(meta)를 읽지 못했습니다.")
    s["history"] = [_public_meta(m, _read_backup(style_id, v)) for v in vs for m in [_read_meta(style_id, v)] if m]
    return s


@router.get("/styles/{style_id}/v/{ver}/view")
def get_view(style_id: str, ver: str):
    _check_style_id(style_id)
    v = _resolve_ver(style_id, ver)
    return FileResponse(os.path.join(_vdir(style_id, v), "view.json"), media_type="application/json",
                        headers={"Cache-Control": "private, max-age=300", "X-Cad-Version": str(v)})


def _disposition(ascii_name: str, orig: Optional[str]) -> str:
    from urllib.parse import quote
    d = f'attachment; filename="{ascii_name}"'
    if orig and orig != ascii_name:
        d += f"; filename*=UTF-8''{quote(orig)}"
    return d


@router.get("/styles/{style_id}/v/{ver}/files/{kind}")
def get_file(style_id: str, ver: str, kind: str, size: Optional[str] = None):
    """kind=zip(전체) · dxf/pdf/ai(?size=S) · aama · rul · allsizes · a4 — 올린 그대로(지문 동일) 내려준다."""
    _check_style_id(style_id)
    v = _resolve_ver(style_id, ver)
    if kind not in FILE_KINDS:
        raise HTTPException(status_code=400, detail="파일 종류는 " + " · ".join(FILE_KINDS) + " 중 하나입니다.")
    vd = _vdir(style_id, v)
    meta = _read_meta(style_id, v) or {}
    if kind == "zip":
        name = f"{style_id}_v{v}.zip"
        return FileResponse(os.path.join(vd, "exports.zip"), media_type="application/zip",
                            headers={"Content-Disposition": _disposition(name, None), "X-Content-SHA256": meta.get("zip_sha256", "")})
    files = [f for f in (meta.get("files") or []) if f.get("kind") == kind]
    if size:
        files = [f for f in files if f.get("size") == size]
    elif len(files) > 1:
        raise HTTPException(status_code=400, detail="사이즈를 골라 주세요(?size=S 처럼). 있는 사이즈: "
                            + ", ".join(str(f.get("size")) for f in files))
    if not files:
        raise HTTPException(status_code=404, detail="그 파일은 이 버전에 없습니다.")
    f = files[0]
    with zipfile.ZipFile(os.path.join(vd, "exports.zip")) as zf:
        if zf.getinfo(f["path"]).file_size > MAX_UNZIPPED_FILE:              # 올릴 때 검사했지만 한 번 더(디스크 사본이 바뀐 경우)
            raise HTTPException(status_code=500, detail="저장된 파일이 비정상적으로 큽니다. 관리자에게 알려 주세요.")
        data = zf.read(f["path"])
    if _sha256(data) != f.get("sha256"):                                    # 내보내기 직전에 한 번 더 — 손상 파일은 내보내지 않는다
        raise HTTPException(status_code=500, detail="저장된 파일 지문이 맞지 않습니다(손상). 관리자에게 알려 주세요.")
    ascii_name = f"{style_id}_v{v}_" + f["path"].rsplit("/", 1)[-1]
    return Response(content=data, media_type=FILE_KINDS[kind][0],
                    headers={"Content-Disposition": _disposition(ascii_name, f.get("orig")), "X-Content-SHA256": f["sha256"]})


@router.post("/styles", dependencies=[Depends(admin_required)])
async def upload_style(request: Request, view: UploadFile = File(...), exports: UploadFile = File(...),
                       note: str = Form("")) -> JSONResponse:
    """🛡 관리자 — view.json + exports.zip 를 새 버전으로 쌓는다(덮어쓰기 없음). 디스크 → 창고 백업 → 되읽기 확인."""
    view_bytes = await view.read(MAX_VIEW_BYTES + 1)
    zip_bytes = await exports.read(MAX_ZIP_BYTES + 1)
    v = validate_upload(view_bytes, zip_bytes)
    code, content = store_version(v, view_bytes, zip_bytes, jwt=_bearer(request), uploaded_by=getattr(request.state, "email", None),
                                  note=note, view_name=view.filename, zip_name=exports.filename)
    return JSONResponse(status_code=code, content=content)


def store_version(v: Dict[str, Any], view_bytes: bytes, zip_bytes: bytes, jwt: Optional[str], uploaded_by: Optional[str],
                  note: str = "", view_name: Optional[str] = None, zip_name: Optional[str] = None,
                  extra: Optional[Dict[str, Any]] = None,
                  extra_files: Optional[Dict[str, bytes]] = None, expect_version: Optional[int] = None) -> Tuple[int, Dict[str, Any]]:
    """expect_version: 이 번호가 아니면(그 사이 다른 저장이 끼었다) 409 — 번호 확인과 쓰기를 같은 자물쇠 안에서(들여오기 v1 · 이름 저장 n+1)
    검사를 통과한(validate_upload) view.json + exports.zip 을 새 버전으로 쌓는다 — 올리기·서버 다시 만들기가 같이 쓴다.
    extra_files: 같이 넣을 파일 — 지금은 «rules.json»(다시 그레이딩한 규칙표)만 받는다(meta 에 지문을 적는다).
    디스크 → 창고 백업 → 되읽기 확인. 반환 (상태 코드, 본문): 201 성공 · 502 창고 실패(디스크에서도 지움)."""
    extra_files = {k: b for k, b in (extra_files or {}).items() if k == "rules.json" and isinstance(b, (bytes, bytearray))}
    style = v["style"]; sid = style["id"]
    restore_missing_with_token(jwt, force=True)                            # 창고에만 있는 옛 버전과 번호가 겹치지 않게
    with _LOCK:
        nums = _versions_on_disk(sid)
        stored = _sb_list(sid + "/", jwt) if (jwt and _sb_config()[0]) else None
        for it in stored or []:
            m = re.fullmatch(r"v([1-9][0-9]{0,5})", (it.get("name") or "") if isinstance(it, dict) else "")
            if m:
                nums.append(int(m.group(1)))
        ver = (max(nums) + 1) if nums else 1
        if expect_version is not None and ver != expect_version:
            return 409, {"detail": (f"«{sid}» 가 그 사이 다른 작업으로 생겼습니다" if expect_version == 1 else
                                    f"그 사이 v{ver - 1} 이 저장됐습니다") + " — 저장하지 않았습니다. 최신 판을 열어 다시 해 주세요.",
                         "style_id": sid, "version": ver}
        safe_extra = {k: v2 for k, v2 in (extra or {}).items() if k in META_EXTRA_KEYS}   # 허용 열쇠만 — 핵심 칸(지문·번호)은 못 덮는다
        meta = {**safe_extra, "style_id": sid, "version": ver, "name": style.get("name"), "garment": style.get("garment"),
                "group": style.get("group"), "sizes": style.get("sizes"), "size_labels": style.get("size_labels"),
                "base_size": style.get("base_size"), "uploaded_at": time.strftime("%Y-%m-%dT%H:%M:%S"),
                "uploaded_by": uploaded_by, "note": (note or "")[:500],
                "view_sha256": _sha256(view_bytes), "zip_sha256": _sha256(zip_bytes),
                "view_bytes": len(view_bytes), "zip_bytes": len(zip_bytes), "source": v.get("source"),
                "files": v.get("files") or [], "view_name": view_name, "zip_name": zip_name}
        if "rules.json" in extra_files:
            meta["rules_sha256"] = _sha256(bytes(extra_files["rules.json"]))  # 복원 때 규칙표가 짝인지 본다
        _write_version(sid, ver, {"view.json": view_bytes, "exports.zip": zip_bytes, **{k: bytes(b) for k, b in extra_files.items()},
                                  "meta.json": json.dumps(meta, ensure_ascii=False, indent=1).encode("utf-8")})
    backup = backup_version(sid, ver, jwt)
    if not backup.get("ok") and not backup.get("skipped"):
        # 창고에 못 넣은 버전은 디스크에서 지운다 — 남겨 두면 «있는 것처럼» 보이다가 재배포 때 조용히 사라진다.
        # 창고는 덮어쓰기·삭제 규칙이 없어 반쯤 올라간 조각은 남는다 → 다음 올리기는 다음 번호를 쓴다.
        # 복원은 meta 가 없거나 지문이 다르면 건너뛰므로 반쪽 조각이 되살아나지 않는다.
        with _LOCK:
            shutil.rmtree(_vdir(sid, ver), ignore_errors=True)
        _log(f"[cad] 창고 백업 실패로 v{ver} 를 디스크에서 지움: {sid}")
        return 502, {
            "detail": f"창고(pattern-cad)에 보관하지 못해 v{ver} 올리기를 취소했습니다. 잠시 후 다시 올려 주세요. "
                      f"(사유: {backup.get('reason') or '알 수 없음'})", "style_id": sid, "version": ver, "backup": backup}
    _log(f"[cad] 새 버전: {sid} v{ver} (view {len(view_bytes):,} B · zip {len(zip_bytes):,} B)")
    return 201, {"style_id": sid, "version": ver, "backup": backup,
                 "view_sha256": meta["view_sha256"], "zip_sha256": meta["zip_sha256"]}


# ══════════════════════════════════════════════════════════════════════
#  F10 2단계 — 서버에서 엔진 돌려 «다시 만들기» (🛡 관리자)
# ══════════════════════════════════════════════════════════════════════
# 왜(비유): 1단계는 stizshop 에서 구워 온 «사진» 만 걸었다. 2단계는 그 «사진기(엔진)» 의 도장 찍힌 사본
#   (vendor/pattern_engine — scripts/sync_pattern_engine.py 가 베낀다)을 서버에 두고 서버가 직접 다시 굽는다.
#   엔진은 «따로 띄운 일꾼 프로세스» 에서 돈다(service.rebuild) — 서버 메모리가 부풀지 않고, 엔진이 죽어도 서버는 산다.
#   재단 패턴 작업은 «한 줄로» 선다(_ENGINE_LOCK) — 일꾼 1명(--workers 1) 서버라 둘이 동시에 돌면 합성 작업까지 느려진다.
# 만든 결과는 «올리기» 와 같은 검사(validate_upload)·같은 저장(store_version: 새 버전·창고 백업·되읽기)을 거친다.
def engine_dir() -> str:
    return os.environ.get("GRADER_PATTERN_ENGINE_DIR") or os.path.join(state.PROJECT_ROOT, "vendor", "pattern_engine")


_ENGINE_LOCK = threading.Lock()                                      # 엔진 작업은 한 번에 하나(한 줄로 세우기)
_QUEUE_LOCK = threading.Lock()
_queue = {"waiting": 0, "marker": 0}                                # 줄 선(시작 전 + 돌고 있는) 작업 수 — 관리자 작업 · 직원 마커 따로
REBUILD_QUEUE_MAX = 3                                                # 이보다 많이 밀리면 429 — 누가 단추를 연타해도 서버가 안 잠긴다
_service_mod: Dict[str, Any] = {}
CHILD_GUARD_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "engine_child")   # 일꾼 안전띠(sitecustomize.py)
# 일꾼 주소 공간 상한(MB, 리눅스만) — 근거(2026-10-19 실측·추정):
#   실제 RAM 최고치는 170~210MB(윈도우 작업 집합), 확정 메모리(commit) 최고 165MB.
#   리눅스 «주소 공간» 은 RAM 이 아니라 예약까지 센다 — 파이썬·numpy·GEOS·MuPDF 라이브러리 지도 + 스레드 스택 + malloc 예약으로
#   검증에서 잰 예약이 약 775MB 였다. 450MB 로 걸면 일꾼이 시작하자마자 MemoryError 로 죽는다.
#   그래서 1024MB(예약 775MB 의 약 1.3배) — «고장으로 끝없이 커질 때» 만 막는 백스톱이다.
#   서버를 지키는 진짜 장치는 oom_score_adj=1000(메모리가 모자라면 커널이 일꾼부터 죽인다).
ENGINE_AS_LIMIT_MB_DEFAULT = 1024
# 일꾼에게 넘기는 환경변수 «허용 목록» — 이것 말고는 하나도 안 넘긴다(API 키·창고 열쇠·GRADER_* 설정이 일꾼에 보이지 않게)
#   2026-10-19 리뷰: «PYTHON*·LC_* 로 시작하면 다 넘김» 은 PYTHONSTARTUP·PYTHONINSPECT·PYTHONPATH(남의 코드 끼워 넣기) 같은 것까지
#   통과시켰다 → 정확한 이름 목록으로. PYTHONPATH 는 넘기지 않고 안전띠 폴더 하나로 새로 만든다(_engine_env).
ENGINE_ENV_ALLOW = {"PATH", "HOME", "LANG", "LANGUAGE", "TZ", "TMP", "TMPDIR", "TEMP", "SYSTEMROOT", "SYSTEMDRIVE", "WINDIR",
                    "COMSPEC", "PATHEXT", "USERPROFILE", "LOCALAPPDATA", "APPDATA", "PROGRAMDATA", "NUMBER_OF_PROCESSORS",
                    "PROCESSOR_ARCHITECTURE",
                    "PYTHONIOENCODING", "PYTHONUTF8", "PYTHONHASHSEED", "PYTHONDONTWRITEBYTECODE", "PYTHONUNBUFFERED",
                    "LC_ALL", "LC_CTYPE", "LC_NUMERIC", "LC_TIME", "LC_COLLATE", "LC_MONETARY", "LC_MESSAGES",
                    "PATTERN_ENGINE_FONT", "PATTERN_ENGINE_CACHE", "PATTERN_ENGINE_NO_SYSTEM_FONTS"}
TOKEN_MIN_LEFT_S = 600                                               # 줄에 넣을 때 로그인이 이만큼은 남아 있어야(엔진 1분 + 줄 3건)
RELOGIN_MSG = "로그인이 끝났거나 곧 끝납니다. 다시 로그인한 뒤 눌러 주세요."
HEAVY_WAIT_S_DEFAULT = 600                                           # 엔진이 «돌던 합성이 끝나기» 를 기다리는 상한(초) — 넘으면 엔진 작업 실패
JOB_KINDS = ("cad-rebuild", "cad-regrade", "cad-version", "cad-marker",
             "cad-import", "cad-names", "cad-intake")                              # 들여오기·이름 확인(F10 6단계)   # GET /jobs 로 볼 수 있는 재단 패턴 작업(마커 = F10 4단계)
PREVIEW_KEEP = 20                                                    # 미리보기 결과 폴더를 최근 몇 개까지 남길지(디스크 절약)
MAX_RULE_CHANGES_META = 300                                          # meta 에 적는 «바뀐 칸» 목록 상한


def load_service():
    """vendor/pattern_engine/service.py 를 불러온다(표준 라이브러리만 쓰는 창구라 불러와도 엔진이 돌지 않는다)."""
    path = os.path.join(engine_dir(), "service.py")
    if not os.path.isfile(path):
        raise HTTPException(status_code=503, detail="서버에 패턴 엔진 사본이 없습니다(vendor/pattern_engine). 관리자에게 알려 주세요.")
    mt = os.path.getmtime(path)
    m = _service_mod.get(path)
    if m is None or m[0] != mt:
        spec = importlib.util.spec_from_file_location("stiz_pattern_engine_service", path)
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        _service_mod[path] = (mt, mod)
    return _service_mod[path][1]


def _engine_env() -> Dict[str, str]:
    """일꾼에게 넘길 환경변수 — 허용 목록만(+ 안전띠·메모리 절약 설정). 엔진은 디스크만 쓰니 열쇠가 필요 없다."""
    env = {k: v for k, v in os.environ.items() if k.upper() in ENGINE_ENV_ALLOW}
    env["PYTHONPATH"] = CHILD_GUARD_DIR                                 # 일꾼이 켜질 때 sitecustomize(안전띠)를 먼저 읽게 — 이것 하나만
    env["GRADER_ENGINE_LIMIT_AS_MB"] = str(os.environ.get("GRADER_ENGINE_LIMIT_AS_MB") or ENGINE_AS_LIMIT_MB_DEFAULT)
    env["GRADER_ENGINE_OOM_ADJ"] = os.environ.get("GRADER_ENGINE_OOM_ADJ", "1000")
    for k in ("OPENBLAS_NUM_THREADS", "OMP_NUM_THREADS", "MKL_NUM_THREADS"):  # 계산 스레드 1개 — 스레드마다 잡는 예약 메모리를 줄인다
        env[k] = "1"
    env["MALLOC_ARENA_MAX"] = "2"                                       # glibc 가 스레드마다 64MB 씩 예약하지 않게
    return env


def _jwt_exp(jwt: Optional[str]) -> Optional[int]:
    """토큰(JWT)에 적힌 만료 시각(초). 서명 검사는 require_auth 가 이미 했다 — 여기서는 «남은 시간» 만 읽는다."""
    try:
        import base64
        part = jwt.split(".")[1]
        data = json.loads(base64.urlsafe_b64decode(part + "=" * (-len(part) % 4)).decode("utf-8"))
        exp = data.get("exp")
        return int(exp) if isinstance(exp, (int, float)) else None
    except Exception:
        return None


def _heavy_wait_s() -> float:
    """엔진이 «돌던 합성이 끝나기» 를 기다리는 상한(초) — GRADER_CAD_HEAVY_WAIT_S 로 바꿀 수 있다(시험용)"""
    try:
        return max(0.05, float(os.environ.get("GRADER_CAD_HEAVY_WAIT_S") or HEAVY_WAIT_S_DEFAULT))
    except ValueError:
        return float(HEAVY_WAIT_S_DEFAULT)


def _short_engine_error(e: Exception) -> str:
    """직원 화면에 보일 짧은 사유 — 엔진 출력 꼬리는 서버 로그에만 남긴다"""
    if isinstance(e, HTTPException):
        return str(e.detail)
    if isinstance(e, TimeoutError):
        return (f"합성 작업이 {max(1, round(_heavy_wait_s() / 60))}분 넘게 끝나지 않아 재단 패턴 작업을 멈췄습니다. "
                "잠시 후 다시 눌러 주세요.")
    t = str(e)
    if "MemoryError" in t:
        return "메모리가 모자라 엔진이 멈췄습니다."
    import re as _re
    m = _re.search(r"종료 코드 (-?\d+)", t)
    if m and int(m.group(1)) < 0:
        return f"엔진이 강제로 멈췄습니다(신호 {-int(m.group(1))} — 메모리 부족일 수 있음)."
    return t.split("\n", 1)[0].replace(" 마지막 출력:", "").strip()[:300] or "알 수 없는 오류"


def _font_name(font: Optional[Dict[str, Any]]) -> Optional[str]:
    if not font:
        return None
    if font.get("source") == "builtin":
        return "Droid Sans Fallback (서버 내장)"
    return os.path.splitext(os.path.basename(font.get("path") or ""))[0] or None


# ── 엔진 작업 공통(다시 만들기·미리보기·새 버전 저장) ─────────────────────────
class JobCancelled(Exception):
    """직원이 «취소» 를 눌렀다(DELETE /api/cad/jobs/{job}) — 실패가 아니라 취소로 끝낸다"""


def _run_engine(job_id: str, svc: Any, call, cancel: Optional[threading.Event] = None) -> Dict[str, Any]:
    """① 엔진끼리 한 줄(_ENGINE_LOCK) ② 돌던 합성이 끝나기를 기다림(상한 _heavy_wait_s — 넘으면 TimeoutError)
    ③ call(on_step, timeout) — 일꾼 프로세스를 띄우는 service 함수. 반환: 엔진 결과(dict)
    cancel: 줄 서 있는 동안 취소되면 엔진을 띄우지 않고 JobCancelled(마커 — 2026-10-20 보안 리뷰)"""
    while not _ENGINE_LOCK.acquire(timeout=0.5):                     # 줄 서는 동안에도 취소를 본다
        if cancel is not None and cancel.is_set():
            raise JobCancelled()
    try:
        if cancel is not None and cancel.is_set():
            raise JobCancelled()
        state.update_job(job_id, stage="wait_heavy", stage_ko="합성 작업이 끝나기를 기다림")
        try:
            gate = state.HEAVY.exclusive_use(timeout=_heavy_wait_s(), cancel=cancel)
            gate.__enter__()                                            # 새 합성은 엔진 뒤로 — 기다리다 포기·취소하면 대기 표시가 풀린다
        except state.Cancelled:
            raise JobCancelled()
        try:
            state.update_job(job_id, stage="start", stage_ko="엔진 시작")
            timeout = float(os.environ.get("GRADER_CAD_ENGINE_TIMEOUT") or 1200)

            def on_step(i: int, n: int, name: str) -> None:
                state.update_job(job_id, done=i - 1, total=n, stage=name, stage_ko=svc.STEP_KO.get(name, name))
            return call(on_step, timeout)
        finally:
            gate.__exit__(None, None, None)                             # 합성 문을 연다(엔진이 실패·취소돼도)
    finally:
        _ENGINE_LOCK.release()


def _engine_run_info(res: Dict[str, Any]) -> Dict[str, Any]:
    """meta·작업 결과에 남길 엔진 실행 기록 — 판·글꼴·메모리(일꾼 최고 RAM · 리눅스 VmPeak/VmHWM)·안전띠·시간"""
    font = res.get("font") or {}
    return {"engine": res.get("engine"), "font": font.get("source"), "font_name": _font_name(font),
            "peak_rss_mb": res.get("peak_rss_mb"), "vm_peak_mb": res.get("vm_peak_mb"), "vm_hwm_mb": res.get("vm_hwm_mb"),
            "guard": res.get("guard"), "timings_s": res.get("timings_s"), "wall_s": res.get("wall_s")}


def _result_style(res: Dict[str, Any], style_id: str) -> Dict[str, Any]:
    """엔진 결과에서 이 스타일 칸 — 없거나 view·zip 파일이 빠졌으면 한국어 오류(KeyError 가 화면에 나가지 않게)"""
    s = (res.get("styles") or {}).get(style_id)
    if not s:
        raise ValueError("엔진 결과에 이 스타일이 없습니다(엔진이 이 스타일을 만들지 못함).")
    for k, ko in (("view", "화면 자료(view.json)"), ("zip", "산출물 묶음(exports.zip)")):
        p = s.get(k)
        if not isinstance(p, str) or not os.path.isfile(p):
            raise ValueError(f"엔진 결과에 {ko}이 없습니다(엔진이 끝까지 만들지 못함).")
    return s


def _store_from_engine(job_id: str, s: Dict[str, Any], style_id: str, jwt: Optional[str], email: Optional[str], note: str,
                       extra: Dict[str, Any], extra_files: Optional[Dict[str, bytes]] = None,
                       expect_version: Optional[int] = None) -> Dict[str, Any]:
    """엔진이 만든 view.json + exports.zip → «올리기» 와 같은 검사·저장(새 버전·창고 백업·되읽기). 실패는 한국어 ValueError."""
    if s.get("checks_bad"):                                                 # 화면에서 다시 잰 길이가 검산.json 과 어긋난 판은 저장하지 않는다
        bad = s["checks_bad"]
        _log(f"[cad] 검산 실패로 저장 안 함 {style_id}: {bad}")
        raise ValueError(f"검산이 맞지 않아 저장하지 않았습니다({len(bad)}칸 · 예: {bad[0][0]} {bad[0][1]}).")
    with open(s["view"], "rb") as f:
        view_bytes = f.read()
    with open(s["zip"], "rb") as f:
        zip_bytes = f.read()
    state.update_job(job_id, stage="store", stage_ko="검사·저장")
    v = validate_upload(view_bytes, zip_bytes)                             # 올리기와 같은 검사 — 서버가 만든 것도 봐주지 않는다
    if v["style"]["id"] != style_id:
        raise ValueError(f"엔진이 다른 스타일을 만들었습니다: {v['style']['id']}")
    code, content = store_version(v, view_bytes, zip_bytes, jwt=jwt, uploaded_by=email, note=note, view_name="view.json",
                                  zip_name="exports.zip", extra=extra, extra_files=extra_files, expect_version=expect_version)
    if code == 409:
        raise ValueError(content.get("detail") or "그 사이 다른 저장이 끼어 저장하지 않았습니다.")
    if code != 201:
        objs = ((content.get("backup") or {}).get("objects") or {}).values()
        expired = (_jwt_exp(jwt) or 1e18) <= time.time()
        if expired or any(str(o.get("reason")) in ("status 401", "status 403") for o in objs):
            raise ValueError("창고에 넣지 못했습니다 — " + RELOGIN_MSG)
        raise ValueError("창고에 넣지 못해 저장을 취소했습니다. 잠시 후 다시 눌러 주세요.")
    return content


def _fail_job(job_id: str, work: str, what: str, e: Exception) -> None:
    """작업 실패 — 화면에는 짧은 한국어 사유, 서버 로그에는 전체 흔적. 작업 폴더는 엔진 출력(engine.log)만 남긴다."""
    _log(f"[cad] {what} 실패:\n{traceback.format_exc()}")                 # 엔진 출력 꼬리는 서버 로그에만
    state.update_job(job_id, status="error", error=f"{what}에 실패했습니다: {_short_engine_error(e)}"[:400])
    try:
        if os.path.isdir(work):
            for nm in os.listdir(work):
                full = os.path.join(work, nm)
                if nm != "engine.log":
                    shutil.rmtree(full, ignore_errors=True) if os.path.isdir(full) else os.remove(full)
    except Exception:
        pass


def _start_engine_job(request: Request, kind: str, total: int, fields: Dict[str, Any], target, args: tuple,
                      need_token: bool = True, lane: str = "cad") -> Tuple[str, int]:
    """엔진 작업을 줄에 세우고 스레드를 띄운다 → (작업 번호, 앞에 선 작업 수). 줄이 가득 차면 429.
    need_token: 창고에 넣을 작업(다시 만들기·새 버전)은 줄 서는 동안 로그인이 끝나면 안 된다 — 남은 시간을 먼저 본다.
    lane: "cad"(관리자 다시 만들기·그레이딩, 3칸) · "marker"(직원 마커, MARKER_QUEUE_MAX 칸) — 칸을 따로 세서
          직원 마커가 줄을 채워 관리자 작업이 429 로 막히는 일이 없게 한다(엔진은 여전히 한 번에 하나)."""
    jwt = _bearer(request)
    key, cap = ("marker", _marker_queue_max()) if lane == "marker" else ("waiting", REBUILD_QUEUE_MAX)
    with _QUEUE_LOCK:
        ahead_now = _queue["waiting"] + _queue["marker"]
    if need_token and auth_required() and jwt:
        exp = _jwt_exp(jwt)
        if exp is not None and exp - time.time() < TOKEN_MIN_LEFT_S + 90 * ahead_now:
            raise HTTPException(status_code=401, detail=RELOGIN_MSG)
    with _QUEUE_LOCK:
        if _queue[key] >= cap:
            if lane == "marker":
                raise HTTPException(status_code=429, detail="다른 마커를 만드는 중입니다(한 번에 하나). 끝난 뒤 다시 눌러 주세요.")
            raise HTTPException(status_code=429, detail=f"재단 패턴 작업이 {_queue['waiting']}건 밀려 있습니다. 끝난 뒤 다시 눌러 주세요.")
        _queue[key] += 1
        ahead = _queue["waiting"] + _queue["marker"] - 1
    try:
        job_id = state.create_job(total=total)
        owner = getattr(request.state, "user_id", None) or getattr(request.state, "email", None)
        state.update_job(job_id, kind=kind, stage="queued", stage_ko="차례 기다림" if ahead else "곧 시작", ahead=ahead,
                         owner=owner, lane=lane, **fields)
        threading.Thread(target=target, daemon=True, args=(job_id, *args)).start()
    except Exception:
        with _QUEUE_LOCK:
            _queue[key] = max(0, _queue[key] - 1)
        raise
    return job_id, ahead


def _queue_done(lane: str = "cad") -> None:
    key = "marker" if lane == "marker" else "waiting"
    with _QUEUE_LOCK:
        _queue[key] = max(0, _queue[key] - 1)


def _rebuild_thread(job_id: str, style_id: str, from_ver: int, group: str, jwt: Optional[str], email: Optional[str], note: str) -> None:
    # 작업 폴더 이름을 web_ 로 시작하게 — 실패해 남은 engine.log 도 기존 30일 자동 정리(state.cleanup)가 치운다
    work = os.path.join(state.get_jobs_dir(), f"web_cad_{time.strftime('%y%m%d_%H%M%S')}_{job_id[:6]}")
    try:
        svc = load_service()
        state.update_job(job_id, out_dir=work)
        res = _run_engine(job_id, svc, lambda on_step, timeout: svc.rebuild(work, groups=[group], on_step=on_step, timeout=timeout,
                                                                          env=_engine_env()))
        s = _result_style(res, style_id)
        run = _engine_run_info(res)
        content = _store_from_engine(job_id, s, style_id, jwt, email, note or f"서버에서 다시 만듦(v{from_ver} 기준)",
                                     extra={"rebuilt_from": from_ver, "engine_run": run, "mesh": s.get("mesh")})
        state.update_job(job_id, status="done", done=len(svc.STEPS), stage="done", stage_ko="끝",
                         result={"style_id": style_id, "version": content["version"], "backup": content.get("backup"), **run})
        _log(f"[cad] 다시 만들기 끝: {style_id} v{from_ver} → v{content['version']} · 일꾼 최고 {res.get('peak_rss_mb')} MB"
             f" · VmPeak {res.get('vm_peak_mb')} MB · 안전띠 {res.get('guard')}")
        shutil.rmtree(work, ignore_errors=True)                             # 결과는 버전 폴더에 들어갔다 — 작업 폴더는 지운다
    except Exception as e:
        _fail_job(job_id, work, "다시 만들기", e)
    finally:
        _queue_done()


@router.post("/styles/{style_id}/v/{ver}/rebuild", dependencies=[Depends(admin_required)])
def rebuild_style(style_id: str, ver: str, request: Request, payload: Optional[Dict[str, Any]] = Body(None)) -> JSONResponse:
    """🛡 관리자 — 서버의 엔진 사본으로 이 스타일을 처음부터 다시 만들어 «새 버전» 으로 쌓는다(옛 버전은 그대로).
    오래 걸리므로(1분 남짓) 작업 번호만 바로 돌려주고, 화면은 GET /api/cad/jobs/{job} 으로 진행률을 묻는다."""
    _check_style_id(style_id)
    v = _resolve_ver(style_id, ver)
    svc = load_service()
    group = getattr(svc, "REBUILDABLE", {}).get(style_id)
    if not group:
        raise HTTPException(status_code=400, detail="이 스타일은 아직 서버에서 다시 만들 수 없습니다(서버에 원본이 없음). stizshop 에서 만들어 올려 주세요.")
    note = (payload or {}).get("note") if isinstance(payload, dict) else None
    note = note[:500] if isinstance(note, str) else ""
    job_id, ahead = _start_engine_job(request, "cad-rebuild", len(svc.STEPS), {"style_id": style_id, "from_version": v}, _rebuild_thread,
                                      (style_id, v, group, _bearer(request), getattr(request.state, "email", None), note))
    return JSONResponse(status_code=202, content={"job_id": job_id, "total": len(svc.STEPS), "ahead": ahead,
                                                  "style_id": style_id, "from_version": v})


# ══════════════════════════════════════════════════════════════════════
#  F10 3단계 — 기준점 끌어서 다시 그레이딩 (규칙표 → 미리보기 → 새 버전)
# ══════════════════════════════════════════════════════════════════════
# 왜(비유): 그레이딩 규칙표는 «사이즈마다 기준점이 얼마나 움직이는가» 를 적은 표다. 화면에서 점을 끌면 표의 칸이 바뀌고,
#   서버의 엔진이 그 표로 전 사이즈를 다시 그린 뒤 맞물림(래글런 몸판=소매·옆선 앞=뒤·소매 밑선·안솔기)을 초록/빨강으로 알려 준다.
#   «저장» 은 늘 새 버전(v n+1) — 옛 버전은 그대로. 공장에 보낸 판(잠금)은 고칠 수 없다.
# 규칙표의 정본: 다시 그레이딩한 버전은 rules.json(넘긴 표 그대로). 그 밖의 버전은 view.json 에서 읽은 표(0.01mm).
def _lock_info(style_id: str, ver: int) -> Optional[Dict[str, Any]]:
    """«공장 보낸 판» 잠금 기록(누가·언제) — 버전 폴더 lock.json(창고 «lock» 의 사본)만 정본.
    meta.json 은 고치지 않는다(창고 meta 는 덮어쓰기 금지라 디스크 meta 를 고치면 둘이 어긋난다 — 검증 지적)."""
    try:
        with open(os.path.join(_vdir(style_id, ver), "lock.json"), encoding="utf-8") as f:
            rec = json.load(f)
        return rec if isinstance(rec, dict) else None
    except Exception:
        return None


def _version_rules(style_id: str, ver: int, group: str, svc: Any) -> Tuple[Dict[str, Any], str]:
    """그 버전의 규칙표(전체) → (표, 출처 "stored"|"view"). 못 읽으면 400(한국어)."""
    try:
        p = os.path.join(_vdir(style_id, ver), "rules.json")
        if os.path.isfile(p):
            with open(p, encoding="utf-8") as f:
                doc = json.load(f)
            return svc.check_rules(group, doc.get("rules") if isinstance(doc, dict) else None, full=True), "stored"
        with open(os.path.join(_vdir(style_id, ver), "view.json"), encoding="utf-8") as f:
            view = json.load(f)
        return svc.check_rules(group, svc.rules_from_view(view), full=True), "view"
    except svc.ServiceError as e:
        raise HTTPException(status_code=400, detail=f"v{ver} 의 그레이딩 규칙을 읽지 못했습니다: {e}")
    except (OSError, ValueError, TypeError, AttributeError, KeyError) as e:
        raise HTTPException(status_code=400, detail=f"v{ver} 의 그레이딩 규칙을 읽지 못했습니다({type(e).__name__}).")


def _group_of(style_id: str, svc: Any) -> Optional[str]:
    return getattr(svc, "REBUILDABLE", {}).get(style_id)


@router.get("/styles/{style_id}/v/{ver}/rules")
def get_rules(style_id: str, ver: str, request: Request) -> Dict[str, Any]:
    """🔒 그 버전의 그레이딩 규칙표 + 화면이 점을 끌 때 쓰는 틀(기준점 → 화면 점 이름·부호) + 잠금·편집 가능 여부"""
    _check_style_id(style_id)
    v = _resolve_ver(style_id, ver)
    svc = load_service()
    group = _group_of(style_id, svc)
    lock = _lock_info(style_id, v)
    out = {"style_id": style_id, "version": v, "group": group, "rules": None, "source": None, "reason": None,
           "sizes": list(getattr(svc, "GRADE_SIZES", [])), "base_size": getattr(svc, "BASE_SIZE", "2XL"), "locked": lock,
           "frame": None, "lock_x": [], "mesh_def": [], "editable": False}
    if not group and _is_imported(style_id, v):                            # 들여온 스타일(F10 6단계) — 규칙표 정본은 묶음 안 rules.json
        out["group"] = "imported"
        try:
            ui = _imported_ui(style_id, v, svc)
        except HTTPException as e:
            out["reason"] = str(e.detail); return out
        out.update(rules=ui["rules"], frame=ui["frame"], sizes=ui["sizes"], base_size=ui["base_size"], source="stored",
                   mesh_def=[{"key": k, "label": lb, "tol": tol} for k, lb, tol in getattr(svc, "IMPORT_MESH_DEF", [])])
        out["editable"] = bool(out["rules"]) and not lock and _can_upload(request)
        return out
    if not group:
        out["reason"] = "이 스타일은 아직 서버에서 다시 그레이딩할 수 없습니다(서버에 원본이 없음)."
        return out
    try:
        out["rules"], out["source"] = _version_rules(style_id, v, group, svc)
    except HTTPException as e:
        out["reason"] = str(e.detail)
    out["frame"] = svc.RULE_FRAME[group]
    out["lock_x"] = [[p, k] for g, p, k in svc.LOCK_X if g == group]
    out["mesh_def"] = [{"key": k, "label": lb, "tol": tol} for k, lb, _s, _n, _t, tol in svc.MESH[group]]
    out["editable"] = bool(out["rules"]) and not lock and _can_upload(request)
    return out


def _parse_rule_request(style_id: str, payload: Any, svc: Any) -> Tuple[int, str, Dict[str, Any], List[Dict[str, Any]]]:
    """요청 {base_version, rules: 바꾼 칸만(최종 규칙 값)} → (기준 버전, 옷, 합친 전체 표, 바뀐 칸 목록). 잘못이면 400·409."""
    if payload is not None and not isinstance(payload, dict):
        raise HTTPException(status_code=400, detail="요청 모양이 다릅니다({base_version, rules}).")
    payload = payload or {}
    bv = payload.get("base_version")
    v = _resolve_ver(style_id, "latest" if bv in (None, "") else str(bv))
    group = _group_of(style_id, svc)
    if not group and _is_imported(style_id, v):                            # 들여온 스타일 — 묶음 rules.json 기준 «바꾼 칸» 표
        if _lock_info(style_id, v):
            raise HTTPException(status_code=409, detail=f"v{v} 은 공장에 보낸 판이라(잠금) 고칠 수 없습니다.")
        ui = _imported_ui(style_id, v, svc)
        try:
            changes = svc.check_imported_changes(ui, payload.get("rules") or {})
        except svc.ServiceError as e:
            raise HTTPException(status_code=400, detail=str(e))
        full = svc.merge_rules(ui["rules"], changes)
        return v, "imported", full, svc.diff_rules(ui["rules"], full)
    if not group:
        raise HTTPException(status_code=400, detail="이 스타일은 아직 서버에서 다시 그레이딩할 수 없습니다(서버에 원본이 없음).")
    if _lock_info(style_id, v):
        raise HTTPException(status_code=409, detail=f"v{v} 은 공장에 보낸 판이라(잠금) 고칠 수 없습니다.")
    base, _src = _version_rules(style_id, v, group, svc)
    try:
        changes = svc.check_rules(group, payload.get("rules") or {})
        full = svc.merge_rules(base, changes)
        svc.check_rules(group, full, full=True)
    except svc.ServiceError as e:
        raise HTTPException(status_code=400, detail=str(e))
    return v, group, full, svc.diff_rules(base, full)


def _prune_previews(keep: int = PREVIEW_KEEP) -> None:
    """미리보기 결과 폴더(web_cadprev_*)는 최근 keep 개만 남긴다 — 끌 때마다 생겨서"""
    try:
        d = state.get_jobs_dir()
        olds = sorted(n for n in os.listdir(d) if n.startswith("web_cadprev_"))
        for n in olds[:-keep] if len(olds) > keep else []:
            shutil.rmtree(os.path.join(d, n), ignore_errors=True)
    except Exception:
        pass


def _mesh_all_green(mesh: Dict[str, Any]) -> bool:
    """저장 조건 — 사이즈가 하나 이상 있고 모든 칸이 ok is True(빨강·판정 못 함 둘 다 안 됨)"""
    sizes = (mesh or {}).get("sizes") or {}
    return bool(sizes) and all(isinstance(rows, list) and rows and all(r.get("ok") is True for r in rows) for rows in sizes.values())


def _rules_hash(table: Dict[str, Any]) -> str:
    """규칙표 지문 — «이 규칙으로 미리보기를 했나» 를 저장 때 맞춰 본다(키 순서와 무관)"""
    return _sha256(json.dumps(table, sort_keys=True, ensure_ascii=False, separators=(",", ":")).encode("utf-8"))


def _first_red(mesh: Dict[str, Any]) -> Optional[str]:
    for s, rows in ((mesh or {}).get("sizes") or {}).items():
        for r in rows:
            if r.get("ok") is False:
                return f"{s} {r.get('label')} {r.get('value')}mm"
    return None


def _regrade_thread(job_id: str, style_id: str, from_ver: int, group: str, table: Dict[str, Any], diff: List[Dict[str, Any]],
                    jwt: Optional[str], email: Optional[str], note: str, preview: bool) -> None:
    prefix = "web_cadprev_" if preview else "web_cad_"
    work = os.path.join(state.get_jobs_dir(), f"{prefix}{time.strftime('%y%m%d_%H%M%S')}_{job_id[:6]}")
    what = "미리보기" if preview else "새 버전 저장"
    try:
        svc = load_service()
        if preview:
            _prune_previews(PREVIEW_KEEP - 1)
        state.update_job(job_id, out_dir=work)
        res = _run_engine(job_id, svc, lambda on_step, timeout: svc.regrade(work, group, table, full=not preview, on_step=on_step,
                                                                          timeout=min(timeout, 900) if preview else timeout,
                                                                          env=_engine_env()))
        s = _result_style(res, style_id)
        mesh = s.get("mesh") or {"ok": None, "sizes": {}}
        run = _engine_run_info(res)
        run["rules_echo_max_mm"] = s.get("rules_echo_max_mm")
        if preview:
            keep = os.path.join(work, "preview_view.json")                 # 화면이 가져갈 그림만 남기고 나머지(중간 결과)는 지운다
            shutil.copyfile(s["view"], keep)
            for nm in os.listdir(work):
                if nm not in ("preview_view.json", "engine.log"):
                    full_p = os.path.join(work, nm)
                    shutil.rmtree(full_p, ignore_errors=True) if os.path.isdir(full_p) else os.remove(full_p)
            state.update_job(job_id, status="done", done=len(svc.STEPS), stage="done", stage_ko="끝", preview_path=keep,
                             result={"preview": True, "style_id": style_id, "base_version": from_ver, "mesh": mesh, "rules": table,
                                     "changes": diff, "checks_bad": s.get("checks_bad") or [],
                                     "view_url": f"/api/cad/jobs/{job_id}/view", **run})
            _log(f"[cad] 미리보기 끝: {style_id} v{from_ver} · 바뀐 칸 {len(diff)} · 맞물림 {'초록' if mesh.get('ok') else '빨강 있음'}"
                 f" · {res.get('wall_s')}초 · 일꾼 최고 {res.get('peak_rss_mb')} MB")
            return
        if not _mesh_all_green(mesh):                                       # 빨강뿐 아니라 «판정 못 함(None)» 칸도 저장하지 않는다
            red = _first_red(mesh)
            raise ValueError(f"맞물림 검산이 전부 초록이 아니라 저장하지 않았습니다(예: {red or '판정 못 한 칸'}). 미리보기로 맞춘 뒤 저장하세요.")
        if _lock_info(style_id, from_ver):                                  # 엔진이 도는 1분 사이 누가 기준 판을 잠갔다 — 잠근 판 위에 쌓지 않는다
            raise ValueError(f"저장하는 사이 v{from_ver} 이 «공장 보냄» 으로 잠겨 저장하지 않았습니다.")
        rules_doc = {"format": "stiz-grade-rules", "version": 1, "style_id": style_id, "group": group,
                     "base_size": svc.BASE_SIZE, "sizes": svc.GRADE_SIZES, "base_version": from_ver,
                     "changes": diff[:MAX_RULE_CHANGES_META], "rules": table}
        content = _store_from_engine(job_id, s, style_id, jwt, email,
                                     note or f"규칙 {len(diff)}칸 바꿔 다시 그레이딩(v{from_ver} 기준)",
                                     extra={"regraded_from": from_ver, "rule_changes": diff[:MAX_RULE_CHANGES_META], "mesh": mesh,
                                            "engine_run": run},
                                     extra_files={"rules.json": json.dumps(rules_doc, ensure_ascii=False, indent=1).encode("utf-8")})
        state.update_job(job_id, status="done", done=len(svc.STEPS), stage="done", stage_ko="끝",
                         result={"style_id": style_id, "version": content["version"], "backup": content.get("backup"),
                                 "base_version": from_ver, "changes": len(diff), "mesh_ok": True, **run})
        _log(f"[cad] 다시 그레이딩 저장: {style_id} v{from_ver} → v{content['version']} · 바뀐 칸 {len(diff)}"
             f" · 일꾼 최고 {res.get('peak_rss_mb')} MB · VmPeak {res.get('vm_peak_mb')} MB")
        shutil.rmtree(work, ignore_errors=True)
    except Exception as e:
        _fail_job(job_id, work, what, e)
    finally:
        _queue_done()


@router.post("/styles/{style_id}/regrade", dependencies=[Depends(admin_required)])
def regrade_preview(style_id: str, request: Request, payload: Optional[Dict[str, Any]] = Body(None)) -> JSONResponse:
    """🛡 관리자 — 규칙 변경(바꾼 칸만) → 전 사이즈 다시 계산 + 맞물림 검산. 저장하지 않는다(약 10초 — 작업 번호로 진행률)."""
    _check_style_id(style_id)
    svc = load_service()
    v, group, full, diff = _parse_rule_request(style_id, payload, svc)
    if group == "imported":                                                 # 들여온 스타일 — 기준 명세 + 규칙표로 13사이즈 다시 만들기(약 7초)
        job_id, ahead = _start_engine_job(request, "cad-regrade", len(svc.RESTYLE_STEPS),
                                          {"style_id": style_id, "from_version": v, "rules_hash": _rules_hash(full), "imported": True},
                                          _restyle_thread, (style_id, v, _changes_of(diff), None, None, None, "", "preview", full, diff),
                                          need_token=False)
        return JSONResponse(status_code=202, content={"job_id": job_id, "total": len(svc.RESTYLE_STEPS), "ahead": ahead,
                                                      "style_id": style_id, "base_version": v, "changes": len(diff)})
    job_id, ahead = _start_engine_job(request, "cad-regrade", len(svc.STEPS),
                                      {"style_id": style_id, "from_version": v, "rules_hash": _rules_hash(full)}, _regrade_thread,
                                      (style_id, v, group, full, diff, None, None, "", True), need_token=False)
    return JSONResponse(status_code=202, content={"job_id": job_id, "total": len(svc.STEPS), "ahead": ahead, "style_id": style_id,
                                                  "base_version": v, "changes": len(diff)})


@router.post("/styles/{style_id}/versions", dependencies=[Depends(admin_required)])
def save_regraded(style_id: str, request: Request, payload: Optional[Dict[str, Any]] = Body(None)) -> JSONResponse:
    """🛡 관리자 — 규칙 변경 + 메모 → 새 버전(v n+1, 옛 버전은 그대로). 맞물림이 빨강이면 저장하지 않는다(1분 남짓)."""
    _check_style_id(style_id)
    svc = load_service()
    v, group, full, diff = _parse_rule_request(style_id, payload, svc)
    if not diff:
        raise HTTPException(status_code=400, detail="바뀐 규칙이 없습니다 — 기준점을 옮기거나 숫자를 바꾼 뒤 저장하세요.")
    # 저장은 «같은 규칙으로 미리보기를 해서 전부 초록» 이었을 때만 — 빨강을 1분 동안 다 만든 뒤에야 거절하지 않게(검증 지적)
    pj = (payload or {}).get("preview_job")
    j = state.get_job(pj) if isinstance(pj, str) and re.fullmatch(r"[0-9a-f]{12}", pj) else None
    if (not j or j.get("kind") != "cad-regrade" or j.get("status") != "done" or j.get("style_id") != style_id
            or j.get("from_version") != v or j.get("rules_hash") != _rules_hash(full)):
        raise HTTPException(status_code=400, detail="이 규칙으로 한 미리보기가 없습니다 — «미리보기» 로 맞물림을 확인한 뒤 저장하세요.")
    if not _mesh_all_green((j.get("result") or {}).get("mesh") or {}):
        raise HTTPException(status_code=400, detail="미리보기에서 맞물림이 전부 초록이 아니어서 저장할 수 없습니다. 빨강 칸을 맞춘 뒤 다시 미리보기 하세요.")
    note = (payload or {}).get("note")
    note = note[:500] if isinstance(note, str) else ""
    if group == "imported":
        job_id, ahead = _start_engine_job(request, "cad-version", len(svc.RESTYLE_STEPS), {"style_id": style_id, "from_version": v},
                                          _restyle_thread, (style_id, v, _changes_of(diff), None, _bearer(request),
                                                            getattr(request.state, "email", None), note, "save", full, diff))
        return JSONResponse(status_code=202, content={"job_id": job_id, "total": len(svc.RESTYLE_STEPS), "ahead": ahead,
                                                      "style_id": style_id, "base_version": v, "changes": len(diff)})
    job_id, ahead = _start_engine_job(request, "cad-version", len(svc.STEPS), {"style_id": style_id, "from_version": v}, _regrade_thread,
                                      (style_id, v, group, full, diff, _bearer(request), getattr(request.state, "email", None), note, False))
    return JSONResponse(status_code=202, content={"job_id": job_id, "total": len(svc.STEPS), "ahead": ahead, "style_id": style_id,
                                                  "base_version": v, "changes": len(diff)})


# ══════════════════════════════════════════════════════════════════════
#  F10 5단계 — 고른 것만 묶어 내려받기 + «공장 보낸 판» 잠금
# ══════════════════════════════════════════════════════════════════════
EXPORT_KINDS = tuple(k for k in FILE_KINDS if k != "zip")
SIZED_KINDS = ("dxf", "pdf", "ai", "spec")                           # spec = 들여온 스타일 사이즈별 명세                                   # 사이즈마다 있는 것 — sizes 로 고를 수 있다
# 고른 것 합계 상한 — 실측(2026-10-20): 한 버전 전부(32개)를 풀면 상의 5.5MB · 하의 17.4MB → 넉넉히 64MB
EXPORT_MAX_TOTAL = 64 * 1024 * 1024


@router.post("/styles/{style_id}/v/{ver}/export")
def export_files(style_id: str, ver: str, payload: Optional[Dict[str, Any]] = Body(None)) -> Response:
    """🔒 형식(kinds)·사이즈(sizes, 선택)를 골라 그것만 ZIP 으로. 파일은 저장된 묶음에서 꺼내 지문을 한 번 더 맞춘 뒤 넣는다."""
    _check_style_id(style_id)
    v = _resolve_ver(style_id, ver)
    p = payload if isinstance(payload, dict) else {}
    kinds, sizes = p.get("kinds"), p.get("sizes")
    if not isinstance(kinds, list) or not kinds or not all(isinstance(k, str) and k in EXPORT_KINDS for k in kinds):
        raise HTTPException(status_code=400, detail="형식(kinds)을 하나 이상 고르세요: " + " · ".join(EXPORT_KINDS))
    if sizes is not None and (not isinstance(sizes, list) or not all(isinstance(s, str) for s in sizes)):
        raise HTTPException(status_code=400, detail="사이즈(sizes)는 이름 목록이어야 합니다.")
    meta = _read_meta(style_id, v) or {}
    picked = [f for f in (meta.get("files") or []) if f.get("kind") in kinds
              and (sizes is None or f.get("kind") not in SIZED_KINDS or f.get("size") in sizes)]
    if not picked:
        raise HTTPException(status_code=404, detail="고른 형식·사이즈의 파일이 이 버전에 없습니다.")
    lock = _lock_info(style_id, v)
    zt = (2026, 1, 1, 0, 0, 0)                                           # ZIP 안 날짜 고정 — 같은 선택이면 같은 묶음(지문 같음)
    # 메모리에 다 올리지 않고 임시 파일에 써서 흘려보낸다(검증 지적) — 보낸 뒤 지운다
    tmp = tempfile.NamedTemporaryFile(prefix="cad_export_", suffix=".zip", delete=False)
    try:
        with tmp, zipfile.ZipFile(os.path.join(_vdir(style_id, v), "exports.zip")) as src,                 zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED, compresslevel=6) as out:
            total = 0
            for f in picked:
                try:
                    info = src.getinfo(f["path"])
                except KeyError:
                    raise HTTPException(status_code=404, detail=f"저장된 묶음에 그 파일이 없습니다: {str(f.get('path'))[:80]}")
                total += info.file_size
                if info.file_size > MAX_UNZIPPED_FILE or total > EXPORT_MAX_TOTAL:
                    raise HTTPException(status_code=413, detail=f"고른 파일이 너무 큽니다(합계 {EXPORT_MAX_TOTAL // 2**20}MB 이하). 형식·사이즈를 줄여 주세요.")
                with src.open(info) as fh:                                  # 파일 하나씩 흘려 넣으며 지문을 잰다(통째로 읽지 않음)
                    h = hashlib.sha256(); zi = zipfile.ZipInfo(f["path"], zt); zi.compress_type = zipfile.ZIP_DEFLATED
                    with out.open(zi, "w") as dst:
                        for chunk in iter(lambda: fh.read(1 << 20), b""):
                            h.update(chunk); dst.write(chunk)
                if h.hexdigest() != f.get("sha256"):
                    raise HTTPException(status_code=500, detail="저장된 파일 지문이 맞지 않습니다(손상). 관리자에게 알려 주세요.")
            man = {"style_id": style_id, "version": v, "kinds": kinds, "sizes": sizes, "from_zip_sha256": meta.get("zip_sha256"),
                   "factory_locked": lock, "files": [{k: f.get(k) for k in ("kind", "size", "path", "orig", "bytes", "sha256")} for f in picked]}
            zi = zipfile.ZipInfo("manifest.json", zt); zi.compress_type = zipfile.ZIP_DEFLATED
            out.writestr(zi, json.dumps(man, ensure_ascii=False, indent=1))
    except BaseException:
        try:
            os.remove(tmp.name)
        except OSError:
            pass
        raise
    name = f"{style_id}_v{v}_{'-'.join(sorted(set(kinds)))}.zip"
    return FileResponse(tmp.name, media_type="application/zip", background=BackgroundTask(_remove_quiet, tmp.name),
                        headers={"Content-Disposition": _disposition(name, None), "X-Cad-Files": str(len(picked))})


def _remove_quiet(path: str) -> None:
    try:
        os.remove(path)
    except OSError:
        pass


def _backup_object(style_id: str, ver: int, disk_name: str, store_name: str, ctype: str, jwt: Optional[str]) -> Dict[str, Any]:
    """버전 폴더의 파일 하나를 창고에 올리고 되읽어 지문을 맞춘다(덮어쓰기 금지 창고 — 이미 있으면 같은 내용인지 본다)"""
    base, _pub = _sb_config()
    if not base:
        return {"ok": False, "skipped": True, "reason": "창고 미설정(로컬)"}
    if not jwt:
        return {"ok": False, "skipped": True, "reason": "로그인 토큰 없음(로컬 무인증)"}
    with open(os.path.join(_vdir(style_id, ver), disk_name), "rb") as f:
        data = f.read()
    key = object_key(style_id, ver, store_name)
    ok, why = _sb_upload(key, data, ctype, jwt)
    back = _sb_download(key, jwt) if ok or why == "status 409" else None
    verified = back is not None and _sha256(back) == _sha256(data)
    return {"ok": verified, "key": key, "uploaded": ok, "verified": verified, "reason": why, "existing": back if (not ok and back) else None}


@router.post("/styles/{style_id}/v/{ver}/lock", dependencies=[Depends(admin_required)])
def lock_version(style_id: str, ver: str, request: Request, payload: Optional[Dict[str, Any]] = Body(None)) -> Dict[str, Any]:
    """🛡 관리자 — «공장 보낸 판» 잠금. 잠근 버전은 다시 그레이딩 저장·미리보기가 안 되고 화면에 «공장 보냄» 이 붙는다.
    기록(누가·언제)은 버전 폴더 lock.json + 창고 «lock» (되읽어 확인) — API 의 meta.locked 로 보인다. 창고에 못 남기면 잠그지 않는다(502)."""
    _check_style_id(style_id)
    v = _resolve_ver(style_id, ver)
    note = (payload or {}).get("note") if isinstance(payload, dict) else None
    note = note[:200] if isinstance(note, str) else ""
    jwt = _bearer(request)
    lp = os.path.join(_vdir(style_id, v), "lock.json")
    with _LOCK:
        if _lock_info(style_id, v):
            raise HTTPException(status_code=409, detail=f"v{v} 은 이미 «공장 보냄» 으로 잠겨 있습니다.")
        meta = _read_meta(style_id, v) or {}
        rec = {"style_id": style_id, "version": v, "locked_by": getattr(request.state, "email", None),
               "locked_at": time.strftime("%Y-%m-%dT%H:%M:%S"), "note": note,
               "zip_sha256": meta.get("zip_sha256"), "view_sha256": meta.get("view_sha256")}
        tmp = lp + f".tmp-{os.getpid()}"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(rec, f, ensure_ascii=False, indent=1)
        os.replace(tmp, lp)
    res = _backup_object(style_id, v, "lock.json", "lock", "application/json", jwt)
    if not res.get("ok") and not res.get("skipped") and res.get("existing"):
        # 창고에 이미 잠금 기록이 있다(지난번 올린 뒤 되읽기만 실패 등) — 그 기록이 이 버전 것이면 그것을 정본으로 받는다
        try:
            old = json.loads(res["existing"].decode("utf-8"))
            if isinstance(old, dict) and old.get("style_id") == style_id and old.get("version") == v:
                rec = old
                with open(lp, "w", encoding="utf-8") as f:
                    json.dump(rec, f, ensure_ascii=False, indent=1)
                res = {"ok": True, "key": res["key"], "adopted": True}
        except Exception:
            pass
    res.pop("existing", None)
    if not res.get("ok") and not res.get("skipped"):
        try:
            os.remove(lp)
        except OSError:
            pass
        raise HTTPException(status_code=502, detail="창고에 잠금 기록을 남기지 못해 잠그지 않았습니다. 잠시 후 다시 눌러 주세요."
                                                    + (" " + RELOGIN_MSG if str(res.get("reason")) in ("status 401", "status 403") else ""))
    # meta.json 은 고치지 않는다 — 화면·API 의 «locked» 는 _public_meta 가 lock.json 에서 읽어 붙인다
    _log(f"[cad] 공장 보냄 잠금: {style_id} v{v} · {rec.get('locked_by')} · 창고 {'확인' if res.get('ok') else '건너뜀'}")
    return {"style_id": style_id, "version": v, "locked": rec, "backup": res}


# ══════════════════════════════════════════════════════════════════════
#  F10 4단계 — 마커·요척 (🔒 로그인 직원 누구나)
# ══════════════════════════════════════════════════════════════════════
# 왜(비유): 마커는 «원단 위에 조각을 테트리스처럼 까는 그림» 이다. 다 깐 길이 ÷ 벌수 = 벌당 요척(원단 m) → 원단 발주·원가.
#   엔진(vendor service.marker)이 일꾼을 두 번 띄운다: ① 원본 → 그레이딩(이 버전과 재단선이 같은지 확인) ② 배치 + PDF·DXF·PNG.
#   배치 계산은 stizshop marker.py 명령줄과 «같은 함수·같은 인자» 라 결과가 같다(시드 고정).
#   무거우니 재단 패턴 작업 줄(_ENGINE_LOCK)·합성과 순서 정하기(HeavyGate exclusive)를 그대로 탄다.
#   결과는 버전에 저장하지 않는다 — 작업 폴더(web_cadmk_*)에만 두고, 30일 자동 정리 + 최근 MARKER_KEEP 개만 남긴다.
#   벌수: 손 입력 표 또는 주문서 엑셀(그레이더 parse_order 재사용 → 피스 규칙: 상의 = 상의 사이즈 칸이 찬 사람 수 · 하의 = 하의 칸).
#   🚫 쇼핑몰 운영 DB 는 읽지 않는다 — 명단은 엑셀 업로드로만 들어온다.
MARKER_KEEP = 30                                                     # 마커 결과 폴더를 최근 몇 개까지 남길지(디스크 절약)
MARKER_ORDER_MAX_BYTES = 10 * 1024 * 1024                            # 주문서 엑셀 상한
MARKER_ORDER_MAX_QTY = 5000                                          # 주문서 행 수량 합 상한(이상한 숫자로 명단이 부풀지 않게)
# 마커 일꾼 주소 공간 상한(MB, 리눅스만 · 운영자 값이 없을 때만) — 근거:
#   ① 입장 판정(admit_marker)이 일꾼 RAM 어림을 상자 512 − 서버 160 − 여유 20 = 332MB 이하로 묶는다(윈도우 실측 맞춤: 상의 24벌 252MB · 하의 15벌 250MB)
#   ② 리눅스 «주소 공간 − RAM» 은 그레이딩 일꾼 실측 약 575MB(라이브러리 지도·스레드·malloc 예약)
#   → 332 + 575 ≈ 907MB, 여유 27% 를 얹어 1152MB. 리눅스에서 마커를 직접 잰 값은 아니다(Render 첫 실행의 결과 memory.marker_vm_peak_mb 로 확인).
#   «고장으로 끝없이 커질 때» 만 막는 백스톱이고, 서버를 지키는 것은 oom_score_adj=1000(일꾼부터 죽는다)과 입장 판정이다.
MARKER_AS_LIMIT_MB_DEFAULT = 1152
MARKER_QUEUE_MAX_DEFAULT = 1                                         # 마커 전용 칸 — 돌고 있는 것 포함 한 번에 하나(GRADER_MARKER_QUEUE_MAX)
MARKER_TIMEOUT_S_DEFAULT = 180                                       # 마커 작업 전체(그레이딩 + 배치 + 파일) 상한 — 넘으면 일꾼을 죽이고 줄·합성 문을 연다
_CANCEL: Dict[str, threading.Event] = {}                             # 작업 번호 → 취소 신호(마커)
_CANCEL_LOCK = threading.Lock()                                      # 취소 창구 ↔ 작업 끝맺음이 엇갈리지 않게(신호가 남지 않게)


def _marker_queue_max() -> int:
    try:
        return max(1, int(os.environ.get("GRADER_MARKER_QUEUE_MAX") or MARKER_QUEUE_MAX_DEFAULT))
    except ValueError:
        return MARKER_QUEUE_MAX_DEFAULT


def _server_mb(svc: Any) -> float:
    """입장 판정의 «서버 몫»(MB) = 지금 이 서버 프로세스의 실제 RAM + 요청 처리 여유 30MB. 못 재면 엔진 기본값(160)"""
    try:
        cur = svc.current_rss_mb()
    except Exception:
        cur = None
    return float(cur) + 30.0 if isinstance(cur, (int, float)) and cur > 0 else float(getattr(svc, "MARKER_SERVER_MB", 160.0))


def _marker_slow(svc: Any) -> float:
    """서버가 엔진 실측 PC 보다 느린 배수(입장 판정의 시간 어림에 곱한다) — GRADER_MARKER_SLOW(0.5~10), 없으면 엔진 기본(1.3 — 2026-10-20 Render 실측 1.10·1.22 + 여유)"""
    try:
        return min(10.0, max(0.5, float(os.environ.get("GRADER_MARKER_SLOW") or svc.MARKER_SLOW)))
    except (ValueError, AttributeError):
        return 1.3


def _marker_timeout_s() -> float:
    """마커 작업 전체 시간 상한(초) — GRADER_MARKER_TIMEOUT_S 로 바꿀 수 있다(30~1200)"""
    try:
        return min(1200.0, max(30.0, float(os.environ.get("GRADER_MARKER_TIMEOUT_S") or MARKER_TIMEOUT_S_DEFAULT)))
    except ValueError:
        return float(MARKER_TIMEOUT_S_DEFAULT)
MARKER_DL = {"pdf": ("application/pdf", "1:1 PDF"), "dxf": ("application/dxf", "DXF(R2000·mm)"),
             "png": ("image/png", "미리보기 그림"), "json": ("application/json", "요약 JSON(조각 자리)")}


def _marker_env() -> Dict[str, str]:
    """마커 일꾼 환경 — 다시 만들기와 같은 허용 목록. 주소 공간 상한은 «운영자 값이 먼저»:
    ① GRADER_MARKER_LIMIT_AS_MB(마커만) ② GRADER_ENGINE_LIMIT_AS_MB(운영자가 엔진 전체에 건 값 — 0 이면 끔) ③ 마커 기본값"""
    env = _engine_env()
    want = (os.environ.get("GRADER_MARKER_LIMIT_AS_MB") or "").strip()
    if want:
        env["GRADER_ENGINE_LIMIT_AS_MB"] = want
    elif not (os.environ.get("GRADER_ENGINE_LIMIT_AS_MB") or "").strip():
        env["GRADER_ENGINE_LIMIT_AS_MB"] = str(MARKER_AS_LIMIT_MB_DEFAULT)
    return env


def order_counts(path: str, svc: Any) -> Dict[str, Any]:
    """주문서 엑셀 → {"top": {사이즈: 벌수}, "bottom": {…}, "rows", "people", "bottom_found", "warnings", "unknown", "sheet"}
    마커 전용 읽기(read_order_members) — 합성이 쓰는 parse_order 는 «상의 사이즈 1차» 규칙이라 하의만 산 줄을 버린다(독립 시험 지적).
    합성 결과를 바꾸지 않으려고 parse_order 는 그대로 두고, 마커는 같은 시트 고르기 규칙으로 상의·하의 열을 따로 읽는다.
    명단 → 벌수는 service.counts_from_members(stizshop marker 와 같은 피스 규칙). 수량이 n 이면 그 줄을 n 벌로 센다."""
    info = read_order_members(path)
    members = info.pop("members")
    return {"top": svc.counts_from_members(members, "top"), "bottom": svc.counts_from_members(members, "bottom"),
            "people": info.get("rows", 0), "pieces_top_bottom": len(members), **info}


# 머리글 이름(공백·nbsp 뺀 글자) — 열 찾기
_HEAD_TOP = ("사이즈", "상의", "상의사이즈", "사이즈(상의)", "상의(사이즈)")
_HEAD_BOT = ("하의", "하의사이즈", "사이즈(하의)", "하의(사이즈)")
_HEAD_QTY = ("수량", "비고", "장수", "매수")
_HEAD_NAME = ("이니셜", "이름")
_HEAD_NUM = ("배번", "번호")
_ORDER_SIZES = ("5XS", "4XS", "3XS", "2XS", "XS", "S", "M", "L", "XL", "2XL", "3XL", "4XL", "5XL")


def _order_size(v: Any, normalize: Any) -> Tuple[str, bool]:
    """칸 하나 → (사이즈, 못 읽음?) — 그레이더 사이즈 규칙(2 XL·L사이즈·10호) + XXL·XXXL 같은 X 여러 개도 받는다"""
    if v is None or str(v).strip() == "":
        return "", False
    s = normalize(v)
    if s:
        return s, False
    t = re.sub(r"(사이즈|SIZE)$", "", str(v).strip().upper().replace(" ", "").replace("-", ""))
    m = re.fullmatch(r"(X{2,5})L", t)
    if m:
        return f"{len(m.group(1))}XL", False
    return "", True


_QTY_UNIT = r"\s*(?:장|벌|개|세트|set|SET|ea|EA)"                    # 수량으로 인정하는 단위 — 숫자 «바로 뒤» 에 붙어야 한다


def _order_qty(v: Any, plain_numbers: bool = False) -> Tuple[int, int, bool]:
    """수량·비고 칸 → (상의 벌수, 하의 벌수, 칸에서 수량을 읽었나).
    숫자 «바로 뒤에 장·벌·개·세트» 가 붙은 것만 수량이다(2026-10-20 독립 시험: «11번으로 변경»·«2024 우승기념»·전화번호·«5XL 요청»
    의 숫자를 벌수로 읽었다). «상의2장 하의1장» 은 따로, 한쪽만 적혀 있으면 그쪽만 그 수 · 다른 쪽은 1(사이즈 칸이 차 있으면 1피스),
    «2장» 처럼 옷 없이 적히면 둘 다 그 수. plain_numbers: 머리글이 «수량·장수·매수» 인 칸 — 맨 숫자(4)도 수량으로 본다(비고는 아님)."""
    if v is None or str(v).strip() == "":
        return 1, 1, False
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if plain_numbers:
            n = max(0, int(v)); return n, n, True
        return 1, 1, False                                              # 비고 칸의 맨 숫자는 수량이 아니다
    t = str(v)
    mt = re.search(r"상의\s*(\d+)" + _QTY_UNIT, t); mb = re.search(r"하의\s*(\d+)" + _QTY_UNIT, t)
    if mt or mb:
        return (int(mt.group(1)) if mt else 1), (int(mb.group(1)) if mb else 1), True
    m = re.search(r"(?<![\d.\-])(\d+)" + _QTY_UNIT, t)
    if m:
        n = int(m.group(1)); return n, n, True
    if plain_numbers and re.fullmatch(r"\s*\d+\s*", t):
        n = int(t); return n, n, True
    return 1, 1, False


def _sheet_members(vals: List[List[Any]], normalize: Any, norm_header: Any, to_str: Any) -> Optional[Dict[str, Any]]:
    """시트 하나 → {"members", "rows", "bottom_found", "skipped", "unknown"} — 머리글(상의/사이즈 열)을 못 찾으면 None"""
    from engine.order import _FOOTER_KEYWORDS
    for hr in range(min(25, len(vals))):
        head = [norm_header(v) for v in vals[hr][:20]]
        c_top = next((i for i, h in enumerate(head) if h in _HEAD_TOP), None)
        c_comb = next((i for i, h in enumerate(head) if "상의" in h and "하의" in h), None)    # «상의|하의» 한 칸에 둘 다
        if c_top is None and c_comb is None:
            continue
        c_bot = next((i for i, h in enumerate(head) if h in _HEAD_BOT), None)
        c_qty = next((i for i, h in enumerate(head) if h in _HEAD_QTY), None)
        plain = c_qty is not None and head[c_qty] in ("수량", "장수", "매수")   # 맨 숫자를 수량으로 보는 칸(비고는 아님)
        c_name = next((i for i, h in enumerate(head) if h in _HEAD_NAME), None)
        c_num = next((i for i, h in enumerate(head) if h in _HEAD_NUM), None)
        start = hr + 1
        if start < len(vals):                                            # 부제 줄(상의 | 하의) — 사이즈 칸 아래
            sub = [norm_header(v) for v in vals[start][:20]]
            if "상의" in sub or "하의" in sub:
                if "상의" in sub: c_top = sub.index("상의")
                if "하의" in sub and c_bot is None: c_bot = sub.index("하의")
                start += 1
        members, rows, skipped, unknown, qty_rows = [], 0, 0, {}, []
        cell = lambda row, c: row[c] if (c is not None and c < len(row)) else None
        for ri, row in enumerate(vals[start:], start=start + 1):
            text = "".join(to_str(v) for v in row[:7]).replace(" ", "")
            if any(k.replace(" ", "") in text for k in _FOOTER_KEYWORDS):
                break                                                    # 배송 정보·연락처 칸 = 명단 끝
            if c_comb is not None and c_top is None:
                parts = re.split(r"[/|,]", to_str(cell(row, c_comb)) or "")
                raw_t, raw_b = parts[0], (parts[1] if len(parts) > 1 else "")
            else:
                raw_t, raw_b = cell(row, c_top), cell(row, c_bot)
            top, bad_t = _order_size(raw_t, normalize)
            bot, bad_b = _order_size(raw_b, normalize)
            for raw, bad in ((raw_t, bad_t), (raw_b, bad_b)):
                if bad:
                    k = to_str(raw)[:12]; unknown[k] = unknown.get(k, 0) + 1
            if not top and not bot:
                if to_str(cell(row, c_name)) or (c_num is not None and to_str(cell(row, c_num))):
                    skipped += 1                                         # 이름·배번은 있는데 사이즈가 비었다 — 알려 준다
                continue                                                 # 빈 줄은 건너뛰고 계속 읽는다(뒤 사람을 버리지 않는다)
            qt, qb, read = _order_qty(cell(row, c_qty), plain)
            if read and ((top and qt != 1) or (bot and qb != 1)):        # 1벌이 아닌 수를 칸에서 읽었다 — 직원이 확인하게 알린다
                qty_rows.append(f"{ri}행 «{to_str(cell(row, c_qty))[:20]}» → " + " · ".join(
                    x for x in ((f"상의 {top} {qt}벌" if top else ""), (f"하의 {bot} {qb}벌" if bot else "")) if x))
            rows += 1
            members += [{"topSize": top, "bottomSize": ""}] * (qt if top else 0)
            members += [{"topSize": "", "bottomSize": bot}] * (qb if bot else 0)
            if len(members) > MARKER_ORDER_MAX_QTY:
                raise HTTPException(status_code=400, detail=f"주문서 수량 합이 {MARKER_ORDER_MAX_QTY}을 넘습니다 — 주문서를 확인해 주세요.")
        if rows:
            return {"members": members, "rows": rows, "bottom_found": c_bot is not None or c_comb is not None,
                    "skipped": skipped, "unknown": unknown, "qty_rows": qty_rows}
    return None


def read_order_members(path: str) -> Dict[str, Any]:
    """마커용 주문서 읽기 — 시트 고르기는 parse_order 와 같은 순서(«주문번호+수량» 표식 시트 → 데이터 있는 첫 시트).
    열기 전에 xlsx_guard(항목별 크기·시트 수 — 엑셀 폭탄)로 막고, 읽기 상한(행 MAX_ROWS·열 MAX_COLS)도 같다.
    파일은 «한 번만» 연다 — 머리글이 없는 사이즈 집계표(양식②)는 이미 읽은 값에 합성과 같은 양식 읽기(상의만)를 쓴다."""
    import openpyxl
    from engine.order import (MAX_ROWS, MAX_COLS, xlsx_guard, _has_order_mark, _normalize_size, _norm_header, _to_str,
                              _parse_form1, _parse_form_stiz, _parse_form2)
    warnings: List[str] = []
    empty = {"members": [], "rows": 0, "bottom_found": False, "warnings": warnings, "unknown": {}, "sheet": None, "skipped": 0}
    why = xlsx_guard(path)
    if why:
        warnings.append(why)
        return empty
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as e:                                               # 서버 경로는 내보내지 않는다
        warnings.append(f"주문서를 열 수 없습니다(엑셀 파일이 아니거나 깨짐: {type(e).__name__}).")
        return empty
    marked = first = None
    fb_marked = fb_first = fb_most = None                                 # 머리글 없는 집계표용(합성과 같은 읽기 · 같은 우선순위)
    try:
        for name in wb.sheetnames:
            try:
                ws = wb[name]
                vals = [list(r) for r in ws.iter_rows(values_only=True, max_row=MAX_ROWS, max_col=MAX_COLS)]
                if (ws.max_row or 0) > MAX_ROWS:
                    warnings.append(f"시트 '{name[:20]}' 가 {MAX_ROWS}행을 넘어 위에서부터 {MAX_ROWS}행까지만 읽었습니다.")
            except Exception as e:
                warnings.append(f"시트 '{name[:20]}' 읽기 실패({type(e).__name__})."); continue
            mark = _has_order_mark(vals)
            got = _sheet_members(vals, _normalize_size, _norm_header, _to_str)
            if got:
                got["sheet"] = name
                if marked is None and mark:
                    marked = got
                if first is None:
                    first = got
                continue
            rows = _parse_form1(vals) or _parse_form_stiz(vals) or _parse_form2(vals)
            if rows:
                if fb_marked is None and mark: fb_marked = rows
                if fb_first is None: fb_first = rows
                if fb_most is None or len(rows) > len(fb_most): fb_most = rows
    finally:
        wb.close()
    best = marked or first
    if best is None:                                                     # 머리글 없는 집계표 — 상의만(파일을 다시 열지 않는다)
        rows = fb_marked or fb_first or fb_most or []
        members = []
        for r in rows:
            try:
                q = max(0, int(str(r.get("qty") or "1")))
            except ValueError:
                q = 1
            members += [{"topSize": r.get("size") or "", "bottomSize": ""}] * q
            if len(members) > MARKER_ORDER_MAX_QTY:
                raise HTTPException(status_code=400, detail=f"주문서 수량 합이 {MARKER_ORDER_MAX_QTY}을 넘습니다 — 주문서를 확인해 주세요.")
        if rows:
            warnings.append("명단 머리글(이니셜·사이즈)이 없는 집계표라 상의 사이즈만 셌습니다 — 하의는 표에 직접 넣어 주세요.")
        else:
            warnings.append("주문서에서 선수/사이즈 정보를 찾지 못했습니다.")
        return {**empty, "members": members, "rows": len(rows), "warnings": warnings}
    if not best["bottom_found"]:
        warnings.append("하의 사이즈 열을 찾지 못했습니다 — 하의 벌수는 표에 직접 넣어 주세요.")
    if best["skipped"]:
        warnings.append(f"이름·배번은 있는데 사이즈가 빈 줄이 {best['skipped']}개 있습니다 — 빠진 사람이 없는지 확인해 주세요.")
    if best["unknown"]:
        warnings.append("사이즈로 못 읽은 칸: " + " · ".join(f"«{k}» {n}개" for k, n in list(best["unknown"].items())[:8]))
    if best.get("qty_rows"):
        q = best["qty_rows"]
        warnings.append(f"수량 칸에서 1벌이 아닌 수를 읽은 줄 {len(q)}개 — 맞는지 확인해 주세요: " + " / ".join(q[:6])
                        + (" …" if len(q) > 6 else ""))
    return {**best, "warnings": warnings}


async def _read_order_upload(up: Any) -> str:
    """올린 주문서 → 디스크(uploads, 7일 정리) 경로. 엑셀이 아니거나 크면 400"""
    name = (getattr(up, "filename", "") or "").lower()
    if not name.endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="주문서는 엑셀(.xlsx) 파일만 올릴 수 있습니다.")
    data = await up.read(MARKER_ORDER_MAX_BYTES + 1)
    if len(data) > MARKER_ORDER_MAX_BYTES:
        raise HTTPException(status_code=400, detail=f"주문서가 너무 큽니다({MARKER_ORDER_MAX_BYTES // 2**20}MB 이하).")
    if not data:
        raise HTTPException(status_code=400, detail="주문서 파일이 비어 있습니다.")
    up.file.seek(0)
    return state.save_upload(up)


@router.post("/order/counts")
async def marker_order_counts(order: UploadFile = File(...)) -> Dict[str, Any]:
    """🔒 주문서 엑셀 → 상의·하의 사이즈별 벌수(피스 규칙). 화면이 표를 채워 직원이 확인한 뒤 마커를 누르게 한다."""
    svc = load_service()
    path = await _read_order_upload(order)
    info = await run_in_threadpool(order_counts, path, svc)              # 엑셀 읽기는 스레드에서 — 서버가 다른 요청을 계속 받게
    if not info["rows"]:
        raise HTTPException(status_code=400, detail=("주문서에서 선수·사이즈를 찾지 못했습니다. " + " ".join(info["warnings"]))[:300])
    return info


def _marker_piece_names(style_id: str, v: int) -> Tuple[List[str], List[str], str]:
    """그 버전 view.json → (사이즈, 조각 이름, view 주소). 못 읽으면 400"""
    vp = os.path.join(_vdir(style_id, v), "view.json")
    try:
        with open(vp, encoding="utf-8") as f:
            st = json.load(f)["style"]
        return list(st["sizes"]), list(st["pieces"]), vp
    except Exception:
        raise HTTPException(status_code=400, detail=f"v{v} 의 화면 자료(view.json)를 읽지 못했습니다.")


def _prune_markers(keep: int = MARKER_KEEP) -> None:
    """마커 결과 폴더(web_cadmk_*)는 최근 keep 개만 — 30일 정리 규칙과 별개로 디스크가 한쪽으로 쌓이지 않게"""
    try:
        d = state.get_jobs_dir()
        olds = sorted(n for n in os.listdir(d) if n.startswith("web_cadmk_"))
        for n in olds[:-keep] if len(olds) > keep else []:
            shutil.rmtree(os.path.join(d, n), ignore_errors=True)
    except Exception:
        pass


def _marker_thread(job_id: str, style_id: str, from_ver: int, group: str, req: Dict[str, Any], rules: Optional[Dict[str, Any]],
                   view_path: str, imp: Optional[Dict[str, Any]] = None) -> None:
    work = os.path.join(state.get_jobs_dir(), f"web_cadmk_{time.strftime('%y%m%d_%H%M%S')}_{job_id[:6]}")
    with _CANCEL_LOCK:
        cancel = _CANCEL.setdefault(job_id, threading.Event())
    try:
        svc = load_service()
        _prune_markers(MARKER_KEEP - 1)
        state.update_job(job_id, out_dir=work)
        limit = _marker_timeout_s()                                     # 전체 상한 — 넘으면 일꾼을 죽이고(=줄·합성 문이 열린다) 실패로
        res = _run_engine(job_id, svc, lambda on_step, timeout: svc.marker(
            work, group, req["counts"], width_cm=req["width_cm"], gap_mm=req["gap_mm"], nap=req["nap"], quick=req["quick"],
            budget=req["budget"], cache_mb=req["cache_mb"], jobs=req["jobs"], seed=req["seed"], exclude=req["exclude"],
            rules=rules, check_view=view_path, timeout=min(timeout, limit), on_step=on_step, env=_marker_env(), cancel=cancel,
            **({"spec_zip": imp["zip"], "spec_files": imp["files"], "title": imp["title"], "per_bundle": req.get("per_bundle"),
                "sizes": req.get("sizes")} if imp else {})),
            cancel=cancel)
        ver = res.get("verify") or {}
        if not ver.get("ok"):                                            # 다시 잰 검산(겹침·간격·폭 밖·식서)이 하나라도 틀리면 내지 않는다
            raise ValueError(f"다시 잰 검산이 맞지 않아 결과를 내지 않았습니다(겹침 {ver.get('overlaps')} · 폭 밖 {ver.get('outside')}"
                             f" · 최소 간격 {ver.get('min_gap_mm')}mm · 식서 {'정상' if ver.get('grain_ok') else '틀림'}).")
        root = os.path.realpath(work)                                   # 결과 주소는 이 작업 폴더 안의 파일만 받는다
        paths = {k: f["path"] for k, f in (res.get("files") or {}).items() if k in MARKER_DL and isinstance(f.get("path"), str)
                 and os.path.isfile(f["path"]) and os.path.commonpath([root, os.path.realpath(f["path"])]) == root}
        for nm in os.listdir(work):                                     # 결과(marker/)·기록만 남기고 중간 결과(그레이딩·화면 자료)는 지운다
            if nm not in ("marker", "engine.log", "engine_graded.log", "result.json", "marker_req.json"):
                full = os.path.join(work, nm)
                shutil.rmtree(full, ignore_errors=True) if os.path.isdir(full) else _remove_quiet(full)
        gr = res.get("graded_run") or {}
        run = _engine_run_info(res)
        mem = {"marker_worker_mb": res.get("peak_rss_mb"), "graded_worker_mb": gr.get("peak_rss_mb"),
               "worker_max_mb": res.get("peak_rss_mb_max"), "marker_vm_peak_mb": res.get("vm_peak_mb"),
               "marker_vm_hwm_mb": res.get("vm_hwm_mb"), "server_peak_mb": svc.peak_rss_mb()}   # 서버 = 이 프로세스(엔진 밖)
        if isinstance(mem["server_peak_mb"], (int, float)) and isinstance(mem["worker_max_mb"], (int, float)):
            mem["sum_mb"] = round(mem["server_peak_mb"] + mem["worker_max_mb"], 1)          # 일꾼은 한 번에 하나라 «서버 + 큰 일꾼» 이 최고
        summ = res.get("summary") or {}
        files = {k: {"bytes": f.get("bytes"), "sha256": f.get("sha256"), "label": MARKER_DL[k][1],
                     "url": f"/api/cad/jobs/{job_id}/files/{k}"} for k, f in (res.get("files") or {}).items() if k in paths}
        state.update_job(job_id, status="done", done=len(svc.MARKER_STEPS), stage="done", stage_ko="끝", marker_files=paths,
                         result={"style_id": style_id, "version": from_ver, "group": group, "request": req,
                                 "summary": summ, "verify": ver, "files": files, "memory": mem,
                                 "view_match": gr.get("view_match"), "rules_used": gr.get("rules_used"),
                                 "graded_sha256": gr.get("graded_sha256"), **run})
        _log(f"[cad] 마커 끝: {style_id} v{from_ver} · {summ.get('length_m')} m · 벌당 {summ.get('yield_m_per_bundle')} m"
             f" · 효율 {summ.get('efficiency_pct')}% · 일꾼 최고 {mem['worker_max_mb']} MB · 서버 {mem['server_peak_mb']} MB"
             f" · VmPeak {mem['marker_vm_peak_mb']} MB · {res.get('wall_s')}초")
    except Exception as e:
        if isinstance(e, JobCancelled) or cancel.is_set():               # 취소 — 실패가 아니다. 만든 것은 버린다
            shutil.rmtree(work, ignore_errors=True)
            state.update_job(job_id, status="error", error="취소했습니다.", cancelled=True)
            _log(f"[cad] 마커 취소됨: {job_id}")
        else:
            _fail_job(job_id, work, "마커", e)
    finally:
        with _CANCEL_LOCK:                                              # 상태(done/error)를 적은 «뒤» 에 지운다 — 취소 창구는 같은 자물쇠 안에서 상태를 본다
            _CANCEL.pop(job_id, None)
        _queue_done("marker")


def _marker_quick(p: Dict[str, Any], request: Request) -> bool:
    """«빠르게»(한 번 깔기)가 서버 기본 — 여러 번 깔기(quick=false)는 관리자만.
    근거(2026-10-20 실측): 일꾼 1명·120초 한도면 여러 번 깔기가 ①② 중간에서 멈춰 상의 8벌은 빠르게와 같은 길이를 4배 시간에 냈다.
    그동안 합성(본업)이 기다리므로, 직원이 누르는 기본 길은 빠르게로 두고 여러 번 깔기는 관리자가 골라서만 쓴다."""
    q = _bool_field(p, "quick", True)
    return True if not _can_upload(request) else q


def _bool_field(p: Dict[str, Any], key: str, default: bool = False) -> bool:
    v = p.get(key, default)
    if not isinstance(v, bool):
        raise HTTPException(status_code=400, detail=f"{key} 는 참/거짓(true/false)이어야 합니다.")
    return v


@router.post("/styles/{style_id}/v/{ver}/marker")
async def marker_start(style_id: str, ver: str, request: Request) -> JSONResponse:
    """🔒 마커·요척 — 원단 폭(cm)·간격(mm)·nap·빠르게(quick)·사이즈별 벌수(counts) 또는 주문서 엑셀 → 작업 번호(202).
    입력: JSON {width_cm, gap_mm, nap, quick, counts: {사이즈: 벌수}, exclude: [조각]}
          또는 multipart: options=<위 JSON 글자> + order=<주문서 .xlsx> (주문서를 주면 그 옷의 벌수를 주문서에서 센다)
    진행률·결과는 GET /api/cad/jobs/{job} · 파일은 GET /api/cad/jobs/{job}/files/{pdf|dxf|png|json}. 버전에는 저장하지 않는다."""
    _check_style_id(style_id)
    v = _resolve_ver(style_id, ver)
    svc = load_service()
    group = _group_of(style_id, svc)
    if not group and _is_imported(style_id, v) and hasattr(svc, "import_style"):
        group = "imported"                                                  # 들여온 스타일 — 그 버전 «명세 묶음» 으로 깐다(F10 6단계)
    if not group or not hasattr(svc, "marker"):
        raise HTTPException(status_code=400, detail="이 스타일은 아직 서버에서 마커를 만들 수 없습니다(서버에 원본이 없음).")
    order_info = None
    if (request.headers.get("content-type") or "").lower().startswith("multipart/form-data"):
        form = await request.form()
        try:
            payload = json.loads(form.get("options") or "{}")
        except (TypeError, ValueError):
            raise HTTPException(status_code=400, detail="options 는 JSON 글자여야 합니다.")
        up = form.get("order")
        if up is not None and hasattr(up, "filename"):
            path = await _read_order_upload(up)
            order_info = await run_in_threadpool(order_counts, path, svc)
            if not order_info["rows"]:
                raise HTTPException(status_code=400, detail=("주문서에서 선수·사이즈를 찾지 못했습니다. " + " ".join(order_info["warnings"]))[:300])
            if not isinstance(payload, dict):
                payload = {}
            side = _order_side(payload, style_id) if group == "imported" else group   # 들여온 스타일 — 주문서의 상의/하의 칸 중 어느 쪽을 셀지
            if side == "bottom" and not order_info["bottom_found"]:
                raise HTTPException(status_code=400, detail="주문서에서 하의 사이즈 열을 찾지 못했습니다 — 하의 벌수를 표에 직접 넣어 주세요.")
            order_info["side"] = side
            # 이 스타일에 없는 사이즈(XS·아동 호수 등)는 빼고 경고로 남긴다 — 주문서 한 줄 때문에 마커 전체가 막히지 않게
            vsizes = set(_marker_piece_names(style_id, v)[0])
            got = order_info[side]
            dropped = {s: n for s, n in got.items() if s not in vsizes}
            payload["counts"] = {s: n for s, n in got.items() if s in vsizes}
            order_info["dropped"] = dropped
            if dropped:
                order_info["warnings"].append("이 스타일에 없는 사이즈라 뺐습니다: " + " · ".join(f"{s} {n}벌" for s, n in dropped.items()))
    else:
        try:
            payload = await request.json()
        except Exception:
            raise HTTPException(status_code=400, detail="요청은 JSON 이어야 합니다({width_cm, gap_mm, counts …}).")
    if not isinstance(payload, dict):
        raise HTTPException(status_code=400, detail="요청 모양이 다릅니다({width_cm, gap_mm, counts …}).")
    sizes, pieces, view_path = _marker_piece_names(style_id, v)
    extra_kw, imp = {}, None
    if group == "imported":                                                 # 한 벌 장수 = 그 버전 조각 장수 합(뺀 조각 빼고) — 입장 판정에 쓴다
        extra_kw["per_bundle"] = _imported_per_bundle(view_path, payload.get("exclude"))
        imp = _imported_marker_input(style_id, v)
    try:
        req = svc.check_marker_request(group, payload.get("counts"), payload.get("width_cm", 150), payload.get("gap_mm", 3),
                                       nap=_bool_field(payload, "nap"), quick=_marker_quick(payload, request),
                                       exclude=payload.get("exclude") or [], sizes=sizes, pieces=pieces, **extra_kw)
        # 예상 메모리·시간(벌수 × 장수 실측식) — 서버 한도를 넘으면 «빠르게» 로 바꾸거나 거절(2026-10-20 보안 리뷰)
        req = svc.admit_marker(req, time_limit_s=_marker_timeout_s(), slow=_marker_slow(svc), server_mb=_server_mb(svc))
    except svc.ServiceError as e:
        raise HTTPException(status_code=400, detail=str(e))
    rules = None
    if os.path.isfile(os.path.join(_vdir(style_id, v), "rules.json")):  # 다시 그레이딩한 버전 — 그 규칙표로 그레이딩해야 그 모양이다
        rules, _src = _version_rules(style_id, v, group, svc)
    total = len(svc.IMARKER_STEPS) if imp else len(svc.MARKER_STEPS)
    job_id, ahead = _start_engine_job(request, "cad-marker", total, {"style_id": style_id, "from_version": v},
                                      _marker_thread, (style_id, v, group, req, rules, view_path, imp), need_token=False, lane="marker")
    return JSONResponse(status_code=202, content={"job_id": job_id, "total": total, "ahead": ahead,
                                                  "style_id": style_id, "version": v, "request": req, "order": order_info})


@router.delete("/jobs/{job_id}")
def cad_job_cancel(job_id: str, request: Request) -> Dict[str, Any]:
    """🔒 마커 작업 취소 — 맡긴 본인 또는 관리자. 줄 서 있으면 엔진을 안 띄우고, 돌고 있으면 일꾼을 죽인다(줄·합성 문이 바로 열린다)."""
    j = state.get_job(job_id) if re.fullmatch(r"[0-9a-f]{12}", job_id or "") else None
    if not j or j.get("kind") != "cad-marker":
        raise HTTPException(status_code=404, detail="취소할 마커 작업이 없습니다.")
    me = getattr(request.state, "user_id", None) or getattr(request.state, "email", None)
    if auth_required() and getattr(request.state, "role", None) != "admin" and j.get("owner") != me:
        raise HTTPException(status_code=403, detail="맡긴 사람이나 관리자만 취소할 수 있습니다.")
    with _CANCEL_LOCK:                                                  # 작업이 막 끝나는 순간과 겹쳐도: 끝난 작업엔 신호를 안 남기고 표시도 안 덮는다
        cur = state.get_job(job_id) or {}
        if cur.get("status") != "running":
            return {"job_id": job_id, "status": cur.get("status"), "cancelled": bool(cur.get("cancelled"))}
        _CANCEL.setdefault(job_id, threading.Event()).set()
        state.update_job_if_running(job_id, stage="cancelling", stage_ko="취소하는 중")
    _log(f"[cad] 마커 취소 요청: {job_id} · {me}")
    return {"job_id": job_id, "status": "cancelling"}


@router.get("/jobs/{job_id}/files/{kind}")
def cad_job_file(job_id: str, kind: str):
    """🔒 마커 결과 파일 — pdf · dxf · png · json (작업이 끝났을 때만, 서버가 다시 켜지면 사라진다 — 다시 만들면 된다)"""
    if kind not in MARKER_DL:
        raise HTTPException(status_code=400, detail="파일 종류는 " + " · ".join(MARKER_DL) + " 중 하나입니다.")
    j = state.get_job(job_id) if re.fullmatch(r"[0-9a-f]{12}", job_id or "") else None
    p = ((j or {}).get("marker_files") or {}).get(kind)
    od = (j or {}).get("out_dir")
    if (not j or j.get("kind") != "cad-marker" or j.get("status") != "done" or not p or not od or not os.path.isfile(p)
            or os.path.commonpath([os.path.realpath(od), os.path.realpath(p)]) != os.path.realpath(od)):
        raise HTTPException(status_code=404, detail="마커 결과가 없습니다(오래됐거나 서버가 다시 켜짐). 마커를 다시 눌러 주세요.")
    r = j.get("result") or {}
    name = f"{r.get('style_id') or 'marker'}_v{r.get('version')}_marker.{kind}"
    ko = None
    if kind != "json":
        g_ko = {"top": "상의", "bottom": "하의"}.get(r.get("group"), "패턴")
        ko = f"마커_{g_ko}_{(r.get('summary') or {}).get('length_m')}m.{kind}"
    return FileResponse(p, media_type=MARKER_DL[kind][0], headers={"Content-Disposition": _disposition(name, ko),
                                                                   "Cache-Control": "no-store"})


# ══════════════════════════════════════════════════════════════════════
#  F10 6단계 — 새 스타일 들여오기 (OptiCad .ai → 조각·기준점 자동 → 사람이 이름 확인)
# ══════════════════════════════════════════════════════════════════════
# 왜(비유): 공장 CAD(OptiCad)가 «이미 그레이딩한» 여러 사이즈 .ai 를 주면, 엔진(stizshop style_import)이 사이즈별 사진을 겹쳐
#   «점마다 한 사이즈 갈 때 얼마나 움직였나» 를 자로 재 규칙표로 만든다. 이름은 자동(조각명_P1…)이라 사람이 확인·고친다.
#   들여온 스타일의 정본 = 묶음(exports.zip) 안 «원본 .ai + 기준 명세(spec_base.json) + 규칙표(rules.json)».
#   이름 고치기·다시 그레이딩·마커는 전부 그 묶음에서 다시 만든다(원본 .ai 를 또 올릴 필요가 없다).
#   올리기·저장은 다른 버전과 같은 길(validate_upload → store_version: 새 버전·창고 백업·되읽기). 원본 .ai 는 {style}/source/ 에도 따로 둔다.
#   🛡 들여오기·이름 저장·다시 그레이딩 = 관리자 · 🔒 확인 목록 보기·마커 = 로그인 직원
IMPORT_TIMEOUT_S_DEFAULT = 600                                       # 들여오기 일꾼 안 시계(원본 읽기·규칙 읽기 상한) — GRADER_CAD_IMPORT_TIMEOUT_S
ZIP_JSON_MAX = 16 * 1024 * 1024                                      # 묶음 안 JSON 하나를 서버가 읽는 상한(rules·confirm 은 실측 0.2MB)
_IMPORT_LOCK = threading.Lock()
_IMPORTING: set = set()                                              # 지금 들여오는 스타일 번호 — 같은 번호 두 번 동시에 못 하게(409)
_ORDER_BOTTOM_RE = re.compile(r"pants|bottom|팬츠|하의|바지|반바지", re.I)


def _import_timeout_s() -> float:
    try:
        return min(1800.0, max(30.0, float(os.environ.get("GRADER_CAD_IMPORT_TIMEOUT_S") or IMPORT_TIMEOUT_S_DEFAULT)))
    except ValueError:
        return float(IMPORT_TIMEOUT_S_DEFAULT)


def _is_imported(style_id: str, ver: int) -> bool:
    """들여온 스타일의 버전인가(meta.imported) — 가스 웜업처럼 «원본에서 다시 만들기» 스타일과 길이 갈린다"""
    return isinstance((_read_meta(style_id, ver) or {}).get("imported"), dict)


def _zip_json(style_id: str, ver: int, name: str) -> Any:
    """그 버전 묶음 안 JSON 하나 — 크기 상한 · meta 의 지문과 맞춰 본 뒤 읽는다(손상 파일로 이름·규칙을 만들지 않게)"""
    want = next((f.get("sha256") for f in ((_read_meta(style_id, ver) or {}).get("files") or []) if f.get("path") == name), None)
    try:
        with zipfile.ZipFile(os.path.join(_vdir(style_id, ver), "exports.zip")) as z:
            info = z.getinfo(name)
            if info.file_size > ZIP_JSON_MAX:
                raise ValueError("너무 큼")
            data = z.read(info)
    except KeyError:
        raise HTTPException(status_code=404, detail=f"이 판 묶음에 {name} 이 없습니다(옛 묶음).")
    except (zipfile.BadZipFile, ValueError, OSError):
        raise HTTPException(status_code=500, detail="저장된 묶음을 읽지 못했습니다. 관리자에게 알려 주세요.")
    if want and _sha256(data) != want:
        raise HTTPException(status_code=500, detail="저장된 파일 지문이 맞지 않습니다(손상). 관리자에게 알려 주세요.")
    try:
        return json.loads(data.decode("utf-8"))
    except (UnicodeDecodeError, ValueError):
        raise HTTPException(status_code=500, detail=f"묶음 안 {name} 을 읽지 못했습니다.")


def _imported_ui(style_id: str, ver: int, svc: Any) -> Dict[str, Any]:
    try:
        return svc.imported_rules_ui(_zip_json(style_id, ver, "rules.json"))
    except svc.ServiceError as e:
        raise HTTPException(status_code=400, detail=f"v{ver} 의 그레이딩 규칙을 읽지 못했습니다: {e}")


def _changes_of(diff: List[Dict[str, Any]]) -> Dict[str, Any]:
    """바뀐 칸 목록 → 일꾼에 넘기는 «바꾼 칸만» 표 {조각: {기준점: {사이즈: [Δx, Δy]}}}"""
    out: Dict[str, Any] = {}
    for d in diff or []:
        out.setdefault(d["piece"], {}).setdefault(d["point"], {})[d["size"]] = list(d["to"])
    return out


def _order_side(payload: Dict[str, Any], style_id: str) -> str:
    """들여온 스타일 마커의 주문서 셈 — 상의 칸(topSize)인가 하의 칸(bottomSize)인가. 화면이 고른 값(order_side) 먼저,
    없으면 스타일 번호·이름에 «팬츠·하의·바지·pants» 가 있으면 하의"""
    s = payload.get("order_side")
    if s in ("top", "bottom"):
        return s
    name = str((_style_summary(style_id) or {}).get("name") or "")
    return "bottom" if _ORDER_BOTTOM_RE.search(style_id + " " + name) else "top"


def _imported_per_bundle(view_path: str, exclude: Any) -> int:
    """한 벌 장수 = 그 버전 조각 장수(qty) 합 − 뺀 조각 — 입장 판정(벌수 × 장수)에 쓴다"""
    ex = {x for x in exclude if isinstance(x, str)} if isinstance(exclude, list) else set()
    try:
        with open(view_path, encoding="utf-8") as f:
            v = json.load(f)
        first = v["style"]["sizes"][0]
        return sum(max(1, int(p.get("qty") or 1)) for p in v["sizes"][first]["pieces"] if p.get("name") not in ex)
    except Exception:
        raise HTTPException(status_code=400, detail="이 판의 화면 자료(view.json)에서 조각 장수를 읽지 못했습니다.")


def _imported_marker_input(style_id: str, ver: int) -> Dict[str, Any]:
    """들여온 스타일 마커 입력 — 그 버전 묶음 + 사이즈별 명세 지문(meta) · 그림 제목"""
    meta = _read_meta(style_id, ver) or {}
    files = {f["path"]: f["sha256"] for f in meta.get("files") or []
             if f.get("kind") == "spec" and isinstance(f.get("path"), str) and isinstance(f.get("sha256"), str)}
    if not files:
        raise HTTPException(status_code=400, detail="이 판에는 사이즈별 명세가 없어 마커를 만들 수 없습니다.")
    return {"zip": os.path.join(_vdir(style_id, ver), "exports.zip"), "files": files,
            "title": f"{str(meta.get('name') or style_id)[:60]} 마커"}


def _store_source(style_id: str, src_path: str, jwt: Optional[str]) -> Dict[str, Any]:
    """원본 .ai 를 {style}/source/original.ai(디스크) + 창고 «{style}/source/original» 에 둔다(되읽어 확인).
    실패해도 들여오기는 성공으로 둔다 — 원본은 버전마다 exports.zip 안(source/original.ai)에도 들어 있어 창고에 이미 백업돼 있다."""
    d = os.path.join(cad_dir(), style_id, "source")
    os.makedirs(d, exist_ok=True)
    dest = os.path.join(d, "original.ai")
    with open(src_path, "rb") as f:
        data = f.read()
    if not os.path.isfile(dest):
        tmp = dest + f".tmp-{os.getpid()}"
        with open(tmp, "wb") as f:
            f.write(data)
        os.replace(tmp, dest)
    base, _ = _sb_config()
    if not base or not jwt:
        res = {"ok": False, "skipped": True, "reason": "창고 미설정(로컬)" if not base else "로그인 토큰 없음(로컬 무인증)"}
    else:
        key = f"{style_id}/source/original"
        if not KEY_RE.fullmatch(key):
            res = {"ok": False, "reason": "창고 열쇠 글자"}
        else:
            ok, why = _sb_upload(key, data, "application/postscript", jwt)
            back = _sb_download(key, jwt) if (ok or why == "status 409") else None
            res = {"ok": back is not None and _sha256(back) == _sha256(data), "key": key, "uploaded": ok, "reason": why}
    res.update(sha256=_sha256(data), bytes=len(data), at=time.strftime("%Y-%m-%dT%H:%M:%S"))
    try:
        with open(os.path.join(d, "backup.json"), "w", encoding="utf-8") as f:
            json.dump(res, f, ensure_ascii=False, indent=1)
    except Exception:
        pass
    _log(f"[cad] 원본 .ai 보관 {'성공' if res.get('ok') else ('건너뜀' if res.get('skipped') else '실패')}: {style_id}")
    return res


def _import_thread(job_id: str, style_id: str, work: str, base: str, name: str, source_name: str, jwt: Optional[str],
                   email: Optional[str], note: str, est: Dict[str, Any]) -> None:
    src = os.path.join(work, "source.ai")
    try:
        svc = load_service()
        state.update_job(job_id, out_dir=work)
        limit = _import_timeout_s()
        res = _run_engine(job_id, svc, lambda on_step, timeout: svc.import_style(
            work, src, base, style_id, name=name or None, source_name=source_name, timeout=min(timeout, limit + 60), max_seconds=limit,
            on_step=on_step, env=_engine_env()))
        s = _result_style(res, style_id)
        if not s.get("check_ok"):                                       # 규칙으로 다시 만든 13사이즈가 원본 CAD 와 안 맞으면 규칙을 믿을 수 없다
            bad = ", ".join((s.get("check_bad") or [])[:3]) or "모양 차"
            raise ValueError(f"원본 대조 검산이 맞지 않아 저장하지 않았습니다({bad}) — 원본 CAD 파일을 확인해 주세요.")
        run = _engine_run_info(res); run["estimate"] = est
        sm = res.get("summary") or {}
        with open(src, "rb") as f:
            src_sha = _sha256(f.read())
        imported = {"base_size": base, "sizes": res.get("sizes"), "pieces": len(res.get("pieces") or []), "source_name": source_name,
                    "source_bytes": os.path.getsize(src), "source_sha256": src_sha, "check_ok": True, "check_max": s.get("check_max"),
                    "confirm_open": s.get("confirm_open"), "seam_knots_total": sm.get("seam_knots_total"), "rules": sm.get("rules")}
        content = _store_from_engine(job_id, s, style_id, jwt, email, note or f"OptiCad 원본에서 들여옴(기준 {base})",
                                     extra={"imported": imported, "engine_run": run}, expect_version=1)   # 같은 번호 재확인은 저장 자물쇠 안에서
        try:
            src_backup = _store_source(style_id, src, jwt)
        except Exception as e:                                          # 원본 따로 보관 실패는 들여오기 실패가 아니다(묶음 안에 원본이 있다)
            src_backup = {"ok": False, "reason": type(e).__name__}
        state.update_job(job_id, status="done", done=len(svc.IMPORT_STEPS), stage="done", stage_ko="끝",
                         result={"style_id": style_id, "version": content["version"], "backup": content.get("backup"),
                                 "source_backup": src_backup, "check_max": s.get("check_max"), "confirm_open": s.get("confirm_open"),
                                 "pieces": imported["pieces"], "sizes": res.get("sizes"), **run})
        _log(f"[cad] 들여오기 끝: {style_id} v{content['version']} · 조각 {imported['pieces']} · 사이즈 {len(res.get('sizes') or [])}"
             f" · 일꾼 최고 {res.get('peak_rss_mb')} MB(어림 {est.get('rss_mb')}) · {res.get('wall_s')}초")
        shutil.rmtree(work, ignore_errors=True)
    except Exception as e:
        _fail_job(job_id, work, "들여오기", e)
    finally:
        _queue_done()
        with _IMPORT_LOCK:
            _IMPORTING.discard(style_id)


@router.post("/styles/import", dependencies=[Depends(admin_required)])
async def import_new_style(request: Request, source: UploadFile = File(...), base_size: str = Form("M"), style_id: str = Form(""),
                           name: str = Form(""), note: str = Form("")) -> JSONResponse:
    """🛡 관리자 — OptiCad «Adobe Export» .ai(여러 사이즈 레이어) → 새 스타일 v1. 작업 번호(202) → GET /api/cad/jobs/{job}.
    검사: 50MB 상한 · 머리(PostScript) · 사이즈 레이어(2개 이상)·기준 사이즈 · 스타일 번호(영문 소문자·숫자·-, 비면 파일 이름에서)
    · 어림 메모리·시간(512MB 상자 안) · 같은 번호가 디스크·창고·들여오는 중이면 409."""
    svc = load_service()
    if not hasattr(svc, "import_style"):
        raise HTTPException(status_code=503, detail="서버 엔진 사본이 들여오기를 아직 모릅니다(엔진 사본을 맞춰 주세요).")
    size = getattr(source, "size", None)                                # 올린 파일은 이미 임시 파일(디스크)에 있다 — 크기만으로 먼저 거른다
    if isinstance(size, int):                                           #   (큰 파일을 서버 메모리에 올리기 전에 · 사이즈 수는 아직 모르니 2로)
        if size > svc.IMPORT_MAX_BYTES:
            raise HTTPException(status_code=400, detail=f"원본이 너무 큽니다({svc.IMPORT_MAX_BYTES // 2 ** 20}MB 이하).")
        try:
            svc.admit_import(size, 2, time_limit_s=_import_timeout_s(), slow=_marker_slow(svc), server_mb=_server_mb(svc))
        except svc.ServiceError as e:
            raise HTTPException(status_code=400, detail=str(e))
    data = await source.read(svc.IMPORT_MAX_BYTES + 1)
    if len(data) > svc.IMPORT_MAX_BYTES:
        raise HTTPException(status_code=400, detail=f"원본이 너무 큽니다({svc.IMPORT_MAX_BYTES // 2 ** 20}MB 이하).")
    if not data:
        raise HTTPException(status_code=400, detail="원본 파일이 비어 있습니다.")
    kind = svc.ai_header_kind(data[:1024])
    if kind == "pdf":
        raise HTTPException(status_code=400, detail="PDF 형식 .ai 는 아직 못 읽습니다 — OptiCad 에서 «Adobe Export»(PostScript) 로 내보내 주세요.")
    if kind != "ps":
        raise HTTPException(status_code=400, detail="OptiCad «Adobe Export» .ai(PostScript) 파일이 아닙니다.")
    if hasattr(svc, "ai_complete") and not svc.ai_complete(data):
        raise HTTPException(status_code=400, detail="파일 끝 표시(%%EOF)가 없습니다 — 내보내기·올리기가 중간에 끊긴 파일 같습니다. OptiCad 에서 다시 내보내 주세요.")
    sizes = await run_in_threadpool(svc.peek_sizes, data)              # 50MB 를 훑는 정규식 — 스레드에서(다른 요청을 막지 않게)
    base = " ".join(str(base_size or "").split()).upper()
    if not sizes:
        raise HTTPException(status_code=400, detail="사이즈 레이어를 찾지 못했습니다 — 여러 사이즈가 레이어로 든 OptiCad 파일인지 확인해 주세요.")
    if base not in sizes:
        raise HTTPException(status_code=400, detail=f"기준 사이즈 {base[:10]} 가 원본에 없습니다 — 있는 사이즈: {' '.join(sizes)}")
    try:
        est = svc.admit_import(len(data), len(sizes), time_limit_s=_import_timeout_s(), slow=_marker_slow(svc), server_mb=_server_mb(svc))
    except svc.ServiceError as e:
        raise HTTPException(status_code=400, detail=str(e))
    sid = " ".join(str(style_id or "").split()).lower() or svc.default_style_id(source.filename or "", data)
    _check_style_id(sid)
    nm = " ".join(str(name or "").split())
    if len(nm) > 80 or any(c in '<>"' or not c.isprintable() for c in nm):
        raise HTTPException(status_code=400, detail="스타일 이름은 80자 이하 · < > \" 없이 써 주세요.")
    src_name = os.path.basename(str(source.filename or "source.ai").replace("\\", "/"))[:120]
    src_name = "".join(c for c in src_name if c.isprintable() and c not in '<>"') or "source.ai"
    note = note[:500] if isinstance(note, str) else ""
    jwt = _bearer(request)
    restore_missing_with_token(jwt)                                    # 재배포 뒤 창고에만 있는 같은 번호를 먼저 디스크로
    if _versions_on_disk(sid):
        raise HTTPException(status_code=409, detail=f"«{sid}» 스타일이 이미 있습니다 — 다른 번호를 쓰거나 그 스타일에서 고쳐 주세요.")
    if _sb_config()[0] and jwt:
        stored = _sb_list(sid + "/", jwt)
        if stored is None:
            raise HTTPException(status_code=502, detail="창고 목록을 읽지 못해 같은 번호가 있는지 모릅니다. 잠시 후 다시 올려 주세요.")
        if any(isinstance(it, dict) and re.fullmatch(r"v[1-9][0-9]{0,5}", it.get("name") or "") for it in stored):
            raise HTTPException(status_code=409, detail=f"«{sid}» 스타일이 창고에 이미 있습니다 — 다른 번호를 써 주세요.")
    with _IMPORT_LOCK:
        if sid in _IMPORTING:
            raise HTTPException(status_code=409, detail=f"«{sid}» 를 지금 들여오는 중입니다.")
        _IMPORTING.add(sid)
    work = os.path.join(state.get_jobs_dir(), f"web_cadimp_{time.strftime('%y%m%d_%H%M%S')}_{hashlib.sha256(os.urandom(8)).hexdigest()[:6]}")
    try:
        os.makedirs(work)
        with open(os.path.join(work, "source.ai"), "wb") as f:          # 영문 이름으로(원래 이름은 기록만) — 경로에 남의 글자가 섞이지 않게
            f.write(data)
        del data
        job_id, ahead = _start_engine_job(request, "cad-import", len(svc.IMPORT_STEPS), {"style_id": sid, "base_size": base},
                                          _import_thread, (sid, work, base, nm, src_name, jwt, getattr(request.state, "email", None),
                                                           note, est))
    except BaseException:
        with _IMPORT_LOCK:
            _IMPORTING.discard(sid)
        shutil.rmtree(work, ignore_errors=True)
        raise
    return JSONResponse(status_code=202, content={"job_id": job_id, "total": len(svc.IMPORT_STEPS), "ahead": ahead, "style_id": sid,
                                                  "base_size": base, "sizes": sizes, "estimate": est})


# ── 패턴 의뢰 → 패턴 자동 제작 + S~3XL 그레이딩(2026-10-22) ─────────────────────
#   의뢰 화면이 제출 직후 부른다. 누가 눌러도 패턴을 만들어 확인표를 의뢰(data.build)에 붙인다.
#   재단 패턴 창고 저장은 관리자 토큰일 때만 — 창고(pattern-cad) 쓰기는 관리자만이라는 규칙을 그대로 둔다.
#   관리자가 누르면 같은 의뢰로 다시 만들어 «검토 대기» 판으로 저장한다(엔진이 7~9초라 다시 만드는 편이 결과를 들고 있는 것보다 단순).
INTAKE_ID_RE = re.compile(r"[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}")
INTAKE_COLS = "id,code,product_name,size_label,base_type,parts,catalog_version,status,data,note,created_by_email,updated_at,deleted_at"
INTAKE_CONFIRM_MAX = 20000                                           # 의뢰에 붙이는 확인표 글자 상한
_INTAKE_LOCK = threading.Lock()
_INTAKE_BUILDING: set = set()                                        # 지금 만드는 의뢰 — 같은 의뢰 두 번 동시에 못 하게(409)


def _intake_rest(method: str, query: str, jwt: Optional[str], body: Any = None, prefer: Optional[str] = None) -> httpx.Response:
    """의뢰 표(pattern_intakes) — 요청자 토큰 그대로(읽기는 모든 직원 · 고치기는 RLS 가 정한 사람만)"""
    base, pub = _sb_config()
    if not base or not jwt:
        raise HTTPException(status_code=503, detail="의뢰 자료를 읽을 수 없습니다(로그인 또는 서버 설정).")
    h = _sb_headers(pub, jwt, "application/json" if body is not None else None)
    if prefer:
        h["Prefer"] = prefer
    return httpx.request(method, f"{base}/rest/v1/pattern_intakes?{query}", headers=h, json=body, timeout=_HTTP_TIMEOUT)


def _intake_read(intake_id: str, jwt: Optional[str]) -> Dict[str, Any]:
    try:
        r = _intake_rest("GET", f"id=eq.{intake_id}&select={INTAKE_COLS}", jwt)
    except httpx.HTTPError:
        raise HTTPException(status_code=502, detail="의뢰 자료를 읽지 못했습니다. 잠시 후 다시 눌러 주세요.")
    if r.status_code in (401, 403):
        raise HTTPException(status_code=401, detail=RELOGIN_MSG)
    if r.status_code != 200:
        raise HTTPException(status_code=502, detail="의뢰 자료를 읽지 못했습니다. 잠시 후 다시 눌러 주세요.")
    rows = r.json()
    if not isinstance(rows, list) or not rows or not isinstance(rows[0], dict):
        raise HTTPException(status_code=404, detail="의뢰를 찾을 수 없습니다.")
    return rows[0]


def _merge_build(old: Any, new: Dict[str, Any]) -> Dict[str, Any]:
    """의뢰 data.build 를 바꿀 때 앞 결과를 잃지 않게 —
    · 실패: 앞에 만든 결과가 있으면 그대로 두고 오류만 붙인다(보이던 패턴이 사라지지 않게)
    · 직원이 다시 만듦(built): 앞서 관리자가 저장한 판은 last_stored 로 남긴다(«재단 패턴 보기» 가 사라지지 않게)"""
    old = old if isinstance(old, dict) else {}
    ok = old.get("status") in ("built", "stored")
    if new.get("status") == "error":
        return {**{k: v for k, v in old.items() if k not in ("error", "error_at")}, "error": new.get("error"), "error_at": new.get("at")} if ok else new
    if new.get("status") == "built":
        last = {k: old.get(k) for k in ("style_id", "ver", "sizes")} if old.get("status") == "stored" else old.get("last_stored")
        if isinstance(last, dict) and last.get("ver"):
            return {**new, "last_stored": last}
    return new


def _intake_save_build(intake_id: str, build: Dict[str, Any], jwt: Optional[str]) -> bool:
    """의뢰 data.build 한 칸만 바꿔 되돌려 보낸다 — 읽은 뒤 누가 고쳤으면(updated_at 다름) 다시 읽어 한 번 더. 성공 여부만(실패해도 작업은 끝)"""
    for _ in range(2):
        try:
            row = _intake_read(intake_id, jwt)
            data = row.get("data") if isinstance(row.get("data"), dict) else {}
            r = _intake_rest("PATCH", f"id=eq.{intake_id}&updated_at=eq.{quote(str(row.get('updated_at')), safe='')}&select=id", jwt,
                             body={"data": {**data, "build": _merge_build(data.get("build"), build)}}, prefer="return=representation")
        except Exception:
            return False
        if r.status_code == 200 and isinstance(r.json(), list) and len(r.json()) == 1:
            return True
        if r.status_code != 200:                                       # 권한 없음(남의 의뢰·검토 끝난 의뢰) 등 — 다시 해도 같다
            return False
    return False


def _intake_thread(job_id: str, intake_id: str, code: str, work: str, admin: bool, jwt: Optional[str], email: Optional[str],
                   sizes: Optional[List[str]] = None) -> None:
    try:
        svc = load_service()
        state.update_job(job_id, out_dir=work)
        res = _run_engine(job_id, svc, lambda on_step, timeout: svc.build_intake(
            work, os.path.join(work, "intake.json"), sizes=sizes or None, timeout=min(timeout, 600), on_step=on_step, env=_engine_env()))
        sid = next(iter(res.get("styles") or {}), "")
        _check_style_id(sid)
        s = _result_style(res, sid)
        sm = res.get("summary") or {}
        try:
            with open(s["confirm"], encoding="utf-8") as f:
                confirm = f.read()[:INTAKE_CONFIRM_MAX]
        except (OSError, KeyError, TypeError):
            confirm = ""
        build = {"status": "built", "style_id": sid, "sizes": res.get("sizes"), "base_size": sm.get("base_size"),
                 "available": sm.get("available") or [], "pieces": len(res.get("pieces") or []),
                 "seams": sm.get("seams"),                             # 솔기 대조 요약 {pairs, ok, bad, warn, issues}
                 "assumed": sm.get("assumed") or [], "at": time.strftime("%Y-%m-%dT%H:%M:%S%z"), "by": email,
                 "engine": (res.get("engine") or {}).get("commit"), "confirm_md": confirm}
        ver = None
        if admin:                                                      # 관리자만 재단 패턴 창고에 «검토 대기» 판으로
            run = _engine_run_info(res)
            restore_missing_with_token(jwt)                            # 재배포 뒤 창고에만 있는 옛 판을 먼저 디스크로(번호가 겹치지 않게)
            have = _versions_on_disk(sid)
            content = _store_from_engine(job_id, s, sid, jwt, email, f"패턴 의뢰 {code} 자동 제작 · 하키티 증분 · 검토 대기",
                                         extra={"intake": {"id": intake_id, "code": code, "review": "pending", "sizes": res.get("sizes"),
                                                           "base_size": sm.get("base_size"),
                                                           "assumed": build["assumed"]}, "engine_run": run},
                                         expect_version=(max(have) + 1) if have else 1)
            ver = content["version"]
            build.update(status="stored", ver=ver)
        saved = _intake_save_build(intake_id, build, jwt)
        with _INTAKE_LOCK:                                             # «끝» 을 보자마자 다시 눌러도 409 가 나지 않게 — 끝 표시보다 먼저 푼다
            _INTAKE_BUILDING.discard(intake_id)
        state.update_job(job_id, status="done", done=len(svc.INTAKE_STEPS), stage="done", stage_ko="끝",
                         result={"style_id": sid, "version": ver, "stored": ver is not None, "saved_to_intake": saved,
                                 "sizes": res.get("sizes"), "base_size": build["base_size"], "available": build["available"],
                                 "pieces": build["pieces"], "assumed": build["assumed"], "seams": build["seams"], "confirm_md": confirm})
        _log(f"[cad] 의뢰 패턴 끝: {code} → {sid} {'v' + str(ver) if ver else '(저장 안 함 · 관리자 아님)'} · 의뢰 기록 {saved}"
             f" · 일꾼 최고 {res.get('peak_rss_mb')} MB · {res.get('wall_s')}초")
        shutil.rmtree(work, ignore_errors=True)
    except Exception as e:
        with _INTAKE_LOCK:                                             # 실패를 보자마자 «다시 만들기» 를 눌러도 되게 — 실패 표시보다 먼저 푼다
            _INTAKE_BUILDING.discard(intake_id)
        _fail_job(job_id, work, "의뢰 패턴 만들기", e)
        err = (state.get_job(job_id) or {}).get("error") or "의뢰 패턴 만들기에 실패했습니다."
        _intake_save_build(intake_id, {"status": "error", "error": err, "at": time.strftime("%Y-%m-%dT%H:%M:%S%z"), "by": email}, jwt)
    finally:
        _queue_done()
        with _INTAKE_LOCK:
            _INTAKE_BUILDING.discard(intake_id)


@router.post("/intake/{intake_id}/build")
def build_from_intake(intake_id: str, request: Request, body: Optional[Dict[str, Any]] = Body(None)) -> JSONResponse:
    """🔒 제출한 패턴 의뢰 → 패턴. 작업 번호(202) → GET /api/cad/jobs/{job}.
    body {"sizes": [...]} 없으면 기준 사이즈(의뢰 «사이즈» 칸)만, 있으면 기준 + 고른 사이즈로 그레이딩(2026-10-22).
    관리자 토큰이면 재단 패턴에 «검토 대기» 새 판으로 저장, 아니면 만들기만 하고 확인표를 의뢰에 붙인다."""
    if not INTAKE_ID_RE.fullmatch(intake_id or ""):
        raise HTTPException(status_code=400, detail="의뢰 번호 형식이 아닙니다.")
    sizes = (body or {}).get("sizes") if isinstance(body, dict) else None
    if sizes is not None and (not isinstance(sizes, list) or len(sizes) > 13
                              or not all(isinstance(x, str) and re.fullmatch(r"[0-9]?X{0,4}[SML]", x) for x in sizes)):
        raise HTTPException(status_code=400, detail="사이즈는 S·M·L·XL·2XL 같은 이름으로 13개까지 골라 주세요.")
    svc = load_service()
    if not hasattr(svc, "build_intake"):
        raise HTTPException(status_code=503, detail="서버 엔진 사본이 의뢰 패턴 만들기를 아직 모릅니다(엔진 사본을 맞춰 주세요).")
    jwt = _bearer(request)
    row = _intake_read(intake_id, jwt)
    if row.get("deleted_at"):
        raise HTTPException(status_code=409, detail="휴지통에 있는 의뢰입니다.")
    if row.get("status") not in ("submitted", "reviewed"):
        raise HTTPException(status_code=409, detail="제출한 의뢰만 패턴을 만들 수 있습니다.")
    admin = (not auth_required()) or getattr(request.state, "role", None) == "admin"
    code = str(row.get("code") or "")[:40]
    with _INTAKE_LOCK:
        if intake_id in _INTAKE_BUILDING:
            raise HTTPException(status_code=409, detail="이 의뢰 패턴을 지금 만드는 중입니다.")
        _INTAKE_BUILDING.add(intake_id)
    work = os.path.join(state.get_jobs_dir(), f"web_cadintake_{time.strftime('%y%m%d_%H%M%S')}_{hashlib.sha256(os.urandom(8)).hexdigest()[:6]}")
    try:
        os.makedirs(work)
        with open(os.path.join(work, "intake.json"), "w", encoding="utf-8") as f:
            json.dump(row, f, ensure_ascii=False)
        job_id, ahead = _start_engine_job(request, "cad-intake", len(svc.INTAKE_STEPS), {"intake_id": intake_id, "intake_code": code,
                                                                                          "sizes": sizes},
                                          _intake_thread, (intake_id, code, work, admin, jwt, getattr(request.state, "email", None), sizes),
                                          need_token=admin)
    except BaseException:
        with _INTAKE_LOCK:
            _INTAKE_BUILDING.discard(intake_id)
        shutil.rmtree(work, ignore_errors=True)
        raise
    return JSONResponse(status_code=202, content={"job_id": job_id, "total": len(svc.INTAKE_STEPS), "ahead": ahead, "code": code,
                                                  "store": admin})


@router.get("/styles/{style_id}/v/{ver}/confirm")
def get_confirm(style_id: str, ver: str, request: Request) -> Dict[str, Any]:
    """🔒 들여온 스타일의 «확인 필요» 목록 — 조각마다 이름·영문·장수·좌우·기준점(자동 이름·자리)·메모·확인함"""
    _check_style_id(style_id)
    v = _resolve_ver(style_id, ver)
    if not _is_imported(style_id, v):
        raise HTTPException(status_code=404, detail="들여온 스타일이 아니라 확인 목록이 없습니다.")
    conf = _zip_json(style_id, v, "confirm.json")
    pieces = [p for p in (conf.get("pieces") if isinstance(conf, dict) else None) or [] if isinstance(p, dict)]
    imp = (_read_meta(style_id, v) or {}).get("imported") or {}
    lock = _lock_info(style_id, v)
    who: Dict[str, Any] = {}                                            # 조각 → 마지막으로 «확인함» 을 켠 사람(옛 판부터 차례로)
    for vv in [x for x in _versions_on_disk(style_id) if x <= v]:
        for c in (_read_meta(style_id, vv) or {}).get("name_changes") or []:
            if isinstance(c, dict) and c.get("field") == "name" and c.get("from") in who:
                who[c.get("to")] = who.pop(c.get("from"))
            if isinstance(c, dict) and c.get("field") == "confirmed":
                if c.get("to"): who[c.get("piece")] = c.get("by")
                else: who.pop(c.get("piece"), None)
    for p in pieces:
        p["confirmed_by"] = who.get(p.get("piece")) if p.get("confirmed") else None
    return {"style_id": style_id, "version": v, "base_size": conf.get("base_size"), "sample": conf.get("sample"), "pieces": pieces,
            "open": sum(1 for p in pieces if not p.get("confirmed")), "total": len(pieces), "locked": lock,
            "editable": not lock and _can_upload(request),
            "import_check": {"ok": imp.get("check_ok"), "max": imp.get("check_max")}}


@router.post("/styles/{style_id}/names", dependencies=[Depends(admin_required)])
def save_names(style_id: str, request: Request, payload: Optional[Dict[str, Any]] = Body(None)) -> JSONResponse:
    """🛡 관리자 — 이름 확인(조각 이름·영문 이름·장수·좌우·기준점 이름·확인함) → 새 버전. 모양은 그대로(일꾼이 점 하나까지 견준다).
    입력: {base_version, pieces: {지금 조각 이름: {name, en_name, qty, mirror, points: {옛: 새}, confirmed}}, note} · 잠긴 판 409"""
    _check_style_id(style_id)
    svc = load_service()
    if not isinstance(payload, dict):
        raise HTTPException(status_code=400, detail="요청 모양이 다릅니다({base_version, pieces}).")
    bv = payload.get("base_version")
    v = _resolve_ver(style_id, "latest" if bv in (None, "") else str(bv))
    if not _is_imported(style_id, v) or not hasattr(svc, "check_names"):
        raise HTTPException(status_code=400, detail="이름 확인은 들여온 스타일만 됩니다.")
    latest = _versions_on_disk(style_id)[-1]
    if v != latest:                                                     # 옛 판 위에 저장하면 그 뒤 판에서 고친 이름이 사라진다
        raise HTTPException(status_code=409, detail=f"v{latest} 이 최신입니다 — 최신 판을 열어 이름을 고쳐 주세요.")
    if _lock_info(style_id, v):
        raise HTTPException(status_code=409, detail=f"v{v} 은 공장에 보낸 판이라(잠금) 고칠 수 없습니다.")
    conf = _zip_json(style_id, v, "confirm.json")
    try:
        clean = svc.check_names(conf, {"pieces": payload.get("pieces")})
    except svc.ServiceError as e:
        raise HTTPException(status_code=400, detail=str(e))
    note = payload.get("note")
    note = note[:500] if isinstance(note, str) else ""
    job_id, ahead = _start_engine_job(request, "cad-names", len(svc.RESTYLE_STEPS), {"style_id": style_id, "from_version": v},
                                      _restyle_thread, (style_id, v, None, clean, _bearer(request), getattr(request.state, "email", None),
                                                        note, "names"))
    return JSONResponse(status_code=202, content={"job_id": job_id, "total": len(svc.RESTYLE_STEPS), "ahead": ahead,
                                                  "style_id": style_id, "base_version": v, "pieces": len(clean["pieces"])})


def _restyle_thread(job_id: str, style_id: str, from_ver: int, rule_changes: Optional[Dict[str, Any]], names: Optional[Dict[str, Any]],
                    jwt: Optional[str], email: Optional[str], note: str, mode: str, full: Optional[Dict[str, Any]] = None,
                    diff: Optional[List[Dict[str, Any]]] = None) -> None:
    """들여온 스타일 다시 묶기 — mode: preview(규칙 미리보기 · 저장 안 함) · save(규칙 새 버전) · names(이름 새 버전)"""
    preview = mode == "preview"
    work = os.path.join(state.get_jobs_dir(), f"{'web_cadprev_' if preview else 'web_cad_'}{time.strftime('%y%m%d_%H%M%S')}_{job_id[:6]}")
    what = {"preview": "미리보기", "save": "새 버전 저장", "names": "이름 저장"}.get(mode, mode)
    try:
        svc = load_service()
        if preview:
            _prune_previews(PREVIEW_KEEP - 1)
        state.update_job(job_id, out_dir=work)
        vd = _vdir(style_id, from_ver)
        res = _run_engine(job_id, svc, lambda on_step, timeout: svc.restyle(
            work, os.path.join(vd, "exports.zip"), rule_changes=rule_changes, names=names, by=email,
            check_view=os.path.join(vd, "view.json") if names else None, timeout=min(timeout, 600), on_step=on_step, env=_engine_env()))
        s = _result_style(res, style_id)
        mesh = s.get("mesh") or {"ok": None, "sizes": {}}
        run = _engine_run_info(res)
        orig = {"ok": s.get("check_ok"), "max": s.get("check_max"), "bad": (s.get("check_bad") or [])[:10]}
        if preview:
            keep = os.path.join(work, "preview_view.json")
            shutil.copyfile(s["view"], keep)
            for nm in os.listdir(work):
                if nm not in ("preview_view.json", "engine.log"):
                    full_p = os.path.join(work, nm)
                    shutil.rmtree(full_p, ignore_errors=True) if os.path.isdir(full_p) else os.remove(full_p)
            state.update_job(job_id, status="done", done=len(svc.RESTYLE_STEPS), stage="done", stage_ko="끝", preview_path=keep,
                             result={"preview": True, "imported": True, "style_id": style_id, "base_version": from_ver, "mesh": mesh,
                                     "rules": full, "changes": diff or [], "checks_bad": [], "import_check": orig,
                                     "view_url": f"/api/cad/jobs/{job_id}/view", **run})
            _log(f"[cad] 미리보기 끝(들여온 스타일): {style_id} v{from_ver} · 바뀐 칸 {len(diff or [])} · 재생성 검산 "
                 f"{'초록' if mesh.get('ok') else '빨강 있음'} · {res.get('wall_s')}초 · 일꾼 최고 {res.get('peak_rss_mb')} MB")
            return
        if not _mesh_all_green(mesh):
            raise ValueError(f"재생성 검산이 전부 초록이 아니라 저장하지 않았습니다(예: {_first_red(mesh) or '판정 못 한 칸'}).")
        if _lock_info(style_id, from_ver):
            raise ValueError(f"저장하는 사이 v{from_ver} 이 «공장 보냄» 으로 잠겨 저장하지 않았습니다.")
        prev = (_read_meta(style_id, from_ver) or {}).get("imported") or {}
        imported = {**prev, "check_ok": bool(s.get("check_ok")), "check_max": s.get("check_max"), "confirm_open": s.get("confirm_open")}
        if mode == "names":
            nch = s.get("name_changes") or []
            extra = {"renamed_from": from_ver, "name_changes": nch[:MAX_RULE_CHANGES_META], "mesh": mesh, "imported": imported, "engine_run": run}
            note = note or f"이름 확인 {len(nch)}건(v{from_ver} 기준)"
        else:
            extra = {"regraded_from": from_ver, "rule_changes": (diff or [])[:MAX_RULE_CHANGES_META], "mesh": mesh, "imported": imported,
                     "engine_run": run}
            note = note or f"규칙 {len(diff or [])}칸 바꿔 다시 그레이딩(v{from_ver} 기준)"
        content = _store_from_engine(job_id, s, style_id, jwt, email, note, extra=extra,
                                     expect_version=(from_ver + 1) if mode == "names" else None)   # 이름: 그 사이 새 판이 생겼으면 409(앞 이름 소실 방지)
        state.update_job(job_id, status="done", done=len(svc.RESTYLE_STEPS), stage="done", stage_ko="끝",
                         result={"style_id": style_id, "version": content["version"], "backup": content.get("backup"),
                                 "base_version": from_ver, "changes": len(diff or s.get("name_changes") or []), "mesh_ok": True,
                                 "confirm_open": s.get("confirm_open"), "import_check": orig, **run})
        _log(f"[cad] {what}(들여온 스타일): {style_id} v{from_ver} → v{content['version']} · 일꾼 최고 {res.get('peak_rss_mb')} MB")
        shutil.rmtree(work, ignore_errors=True)
    except Exception as e:
        _fail_job(job_id, work, what, e)
    finally:
        _queue_done()


# ══════════════════════════════════════════════════════════════════════
#  작업 진행률
# ══════════════════════════════════════════════════════════════════════
@router.get("/jobs/{job_id}")
def cad_job(job_id: str) -> Dict[str, Any]:
    """🔒 재단 패턴 작업 진행률 — running/done/error · 단계(한국어) · 끝나면 새 버전 번호(또는 미리보기 결과)."""
    j = state.get_job(job_id) if re.fullmatch(r"[0-9a-f]{12}", job_id or "") else None
    if not j or j.get("kind") not in JOB_KINDS:
        raise HTTPException(status_code=404, detail="작업을 찾을 수 없습니다(서버 재시작 등). 다시 눌러 주세요.")
    return {k: j.get(k) for k in ("kind", "status", "done", "total", "stage", "stage_ko", "error", "result", "style_id",
                                  "from_version", "ahead", "cancelled")}


@router.get("/jobs/{job_id}/view")
def cad_job_view(job_id: str):
    """🔒 미리보기로 다시 그린 view.json — 저장하지 않은 그림(작업이 끝났을 때만)"""
    j = state.get_job(job_id) if re.fullmatch(r"[0-9a-f]{12}", job_id or "") else None
    p = (j or {}).get("preview_path")
    if not j or j.get("kind") != "cad-regrade" or j.get("status") != "done" or not p or not os.path.isfile(p):
        raise HTTPException(status_code=404, detail="미리보기 그림이 없습니다(오래됐거나 서버가 다시 켜짐). 미리보기를 다시 눌러 주세요.")
    return FileResponse(p, media_type="application/json", headers={"Cache-Control": "no-store"})
